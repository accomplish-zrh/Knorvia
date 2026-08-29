"""Structured Office authoring tool for the chat agent.

``office_document`` lets the model create and edit ``.xlsx`` / ``.docx`` /
``.pptx`` files through a JSON action schema instead of writing openpyxl /
python-docx / python-pptx code. Typical flow:

    create → add_sheet / write_cells → formula → style → chart → read
    → (optional) export_doc / export_slide / screenshot_hint

All file paths are resolved against the turn workspace (``_workspace_dir`` or
``_sandbox_workdir``). When neither is injected the tool falls back to CWD and
reports the absolute path in the result. Failures return
``ToolResult(success=False)`` and never propagate.
"""

from __future__ import annotations

import csv
from io import StringIO
import json
from pathlib import Path
import re
from typing import Any

from knorvia.core.tool_protocol import BaseTool, ToolDefinition, ToolResult
from knorvia.services.office_draft import DraftError, OfficeDraftStore
from knorvia.services.univer_container import (
    UNIT_TYPES,
    ContainerError,
    add_unit,
    export_unit,
    list_units,
    open_univer,
    pack_univer,
)
from knorvia.tools.prompting import load_prompt_hints

WRITE_ACTIONS = (
    "create",
    "add_sheet",
    "write_cells",
    "formula",
    "style",
    "chart",
    "read",
    "export_doc",
    "export_slide",
    "screenshot_hint",
    "container_new",
    "container_add",
    "container_export_unit",
)
LIFECYCLE_ACTIONS = ("ready", "merge", "discard", "status")
DRAFT_ACTIONS = ("create",) + LIFECYCLE_ACTIONS
ACTIONS = WRITE_ACTIONS + LIFECYCLE_ACTIONS
CONTAINER_ACTIONS = ("container_new", "container_add", "container_export_unit")

_SHEET_FORBIDDEN = re.compile(r"[:\\/?*\[\]]")
_HEADING = re.compile(r"^(#{1,3})\s+(.*\S)\s*$")
_UL_ITEM = re.compile(r"^[-*]\s+(.*)$")
_TABLE_SEP = re.compile(r"^\|?\s*:?-+:?\s*(\|\s*:?-+:?\s*)+\|?\s*$")
_A1 = re.compile(r"^[A-Za-z]+\d+$")

_TOOL_DESCRIPTION = (
    "Create and edit structured Office files (xlsx / docx / pptx) without "
    "writing spreadsheet code. Writes default to an isolated draft "
    "(`as_draft=true`); the user previews a review card and confirms before "
    "files land in the turn workspace. Typical flow: create → add_sheet/"
    "write_cells → formula → style → chart → read (self-check) → optional "
    "export_doc / export_slide → draft_action=ready. `create` starts a "
    'workbook (`file` required). `write_cells` accepts {"A1": "Title", '
    '"B2": 123} or a 2D array. `formula` writes formula strings such as '
    '{"D2": "=SUM(B2:C2)"}. `style` takes [{target:"A1:D1", bold:true, '
    'bg:"#B0501E", color:"#FFFFFF", font_size:12}]. `chart` is '
    '{type:"bar|line|pie", data_range:"A1:B5", title:"..."}. `read` '
    "returns CSV of a range so you can verify writes. `export_doc` renders a "
    "Markdown subset (#/##/### headings, - lists, |a|b| tables, paragraphs) "
    "to .docx. `export_slide` builds a 16:9 deck from an outline (level-1 "
    "heading = new slide title; bullets = body). `screenshot_hint` returns a "
    "text grid (and a PNG preview when possible). Multi-unit `.univer` ZIP "
    "containers: `container_new` → `container_add` (unit_type=sheet|doc|slide, "
    "optional source_file / refs) → `container_export_unit`. Independent "
    "draft actions: ready | merge | discard | status (also via `draft_action`)."
)


def _fail(message: str, **meta: Any) -> ToolResult:
    return ToolResult(content=message, success=False, metadata=meta)


def _ok(message: str, output_file: str, **meta: Any) -> ToolResult:
    return ToolResult(
        content=message,
        success=True,
        metadata={"output_file": output_file, **meta},
    )


def _maybe_json(value: Any) -> Any:
    if not isinstance(value, str):
        return value
    text = value.strip()
    if not text:
        return value
    if text[0] not in "{[":
        return value
    try:
        return json.loads(text)
    except json.JSONDecodeError:
        return value


def _as_bool(value: Any, default: bool = False) -> bool:
    if isinstance(value, bool):
        return value
    if isinstance(value, str):
        return value.strip().lower() in {"1", "true", "yes", "on"}
    if value is None:
        return default
    return bool(value)


def _workspace(kwargs: dict[str, Any]) -> tuple[Path, bool]:
    raw = str(kwargs.get("_workspace_dir") or kwargs.get("_sandbox_workdir") or "").strip()
    if raw:
        path = Path(raw).expanduser().resolve()
        path.mkdir(parents=True, exist_ok=True)
        return path, False
    return Path.cwd().resolve(), True


def _resolve_file(workspace: Path, file_name: str) -> Path:
    raw = Path(str(file_name or "").strip())
    if not str(file_name or "").strip() or raw.name in {".", ".."}:
        raise ValueError("`file` is required and must be a relative file name.")
    candidate = raw if raw.is_absolute() else workspace / raw
    resolved = candidate.expanduser().resolve()
    try:
        resolved.relative_to(workspace)
    except ValueError as exc:
        raise ValueError(f"path is outside the turn workspace: {file_name}") from exc
    return resolved


def _rel(path: Path, workspace: Path) -> str:
    try:
        return path.resolve().relative_to(workspace).as_posix()
    except ValueError:
        return str(path)


def _ensure_suffix(path: Path, suffix: str) -> Path:
    if path.suffix.lower() != suffix.lower():
        return path.with_suffix(suffix)
    return path


def _sheet_name(raw: Any, default: str = "Sheet") -> str:
    name = str(raw or "").strip() or default
    if _SHEET_FORBIDDEN.search(name) or name.startswith("'"):
        raise ValueError(f"invalid sheet name {name!r}")
    if len(name) > 31:
        raise ValueError("sheet name must be 31 characters or fewer")
    return name


def _load_openpyxl():
    try:
        from openpyxl import Workbook, load_workbook
        from openpyxl.styles import Font, PatternFill
        from openpyxl.utils import get_column_letter, range_boundaries
    except ImportError as exc:  # pragma: no cover - env without documents extra
        raise ValueError(
            "openpyxl is not installed; install knorvia with the documents extra."
        ) from exc
    return Workbook, load_workbook, Font, PatternFill, get_column_letter, range_boundaries


def _open_workbook(path: Path):
    _Workbook, load_workbook, *_rest = _load_openpyxl()
    if not path.is_file():
        raise ValueError(f"Workbook {path.name!r} does not exist. Call action=create first.")
    return load_workbook(path)


def _save_workbook(workbook: Any, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


def _pick_sheet(workbook: Any, sheet: str | None, *, create_missing: bool = False) -> Any:
    """Resolve the target worksheet.

    ``create_missing`` (used by write-path actions) auto-creates a named
    sheet when absent — the agent should not have to issue a separate
    ``add_sheet`` call before the first ``write_cells``.
    """
    if sheet:
        name = _sheet_name(sheet, default=sheet)
        if name not in workbook.sheetnames:
            if create_missing:
                return workbook.create_sheet(title=name)
            raise ValueError(
                f"Sheet {name!r} not found. Existing: {', '.join(workbook.sheetnames)}"
            )
        return workbook[name]
    return workbook.active


def _rgb(value: str) -> str:
    raw = str(value or "").strip().lstrip("#")
    if len(raw) == 6 and all(ch in "0123456789abcdefABCDEF" for ch in raw):
        return "FF" + raw.upper()
    if len(raw) == 8 and all(ch in "0123456789abcdefABCDEF" for ch in raw):
        return raw.upper()
    raise ValueError(f"invalid colour {value!r}; use #RRGGBB")


def _write_mapping(ws: Any, mapping: dict[Any, Any]) -> tuple[int, int]:
    written = 0
    formulas = 0
    for key, value in mapping.items():
        addr = str(key).strip().upper()
        if not _A1.match(addr):
            raise ValueError(f"invalid cell address {key!r}; use A1-style keys")
        if isinstance(value, str) and value.lstrip().startswith("="):
            formulas += 1
        ws[addr] = value
        written += 1
    return written, formulas


def _write_grid(ws: Any, rows: list[Any], start: str = "A1") -> tuple[int, int]:
    _Workbook, _load, _Font, _Fill, get_column_letter, range_boundaries = _load_openpyxl()
    min_col, min_row, _max_col, _max_row = range_boundaries(start)
    written = 0
    formulas = 0
    for r_offset, row in enumerate(rows):
        if not isinstance(row, (list, tuple)):
            raise ValueError("cells 2D array must be a list of rows")
        for c_offset, value in enumerate(row):
            addr = f"{get_column_letter(min_col + c_offset)}{min_row + r_offset}"
            if isinstance(value, str) and value.lstrip().startswith("="):
                formulas += 1
            ws[addr] = value
            written += 1
    return written, formulas


def _apply_cells(ws: Any, cells: Any) -> tuple[int, int]:
    payload = _maybe_json(cells)
    if isinstance(payload, dict):
        return _write_mapping(ws, payload)
    if isinstance(payload, list):
        return _write_grid(ws, payload)
    raise ValueError('`cells` must be an object {"A1": value, ...} or a 2D array [["a","b"],[1,2]]')


def _action_create(
    path: Path, kwargs: dict[str, Any], workspace: Path, cwd_fallback: bool
) -> ToolResult:
    Workbook, *_rest = _load_openpyxl()
    path = _ensure_suffix(path, ".xlsx")
    workbook = Workbook()
    ws = workbook.active
    ws.title = _sheet_name(kwargs.get("sheet"), default="Sheet")
    _save_workbook(workbook, path)
    rel = _rel(path, workspace)
    extra = f" (absolute {path})" if cwd_fallback else ""
    return _ok(f"Created workbook {rel} with sheet {ws.title!r}{extra}.", rel, sheet=ws.title)


def _action_add_sheet(path: Path, kwargs: dict[str, Any], workspace: Path) -> ToolResult:
    workbook = _open_workbook(path)
    name = _sheet_name(kwargs.get("sheet"), default="")
    if not str(kwargs.get("sheet") or "").strip():
        raise ValueError("`sheet` is required for add_sheet")
    if name in workbook.sheetnames:
        raise ValueError(f"Sheet {name!r} already exists")
    workbook.create_sheet(title=name)
    _save_workbook(workbook, path)
    rel = _rel(path, workspace)
    return _ok(f"Added sheet {name!r} to {rel}.", rel, sheet=name)


def _action_write_cells(path: Path, kwargs: dict[str, Any], workspace: Path) -> ToolResult:
    if kwargs.get("cells") in (None, "", [], {}):
        raise ValueError("`cells` is required for write_cells")
    workbook = _open_workbook(path)
    ws = _pick_sheet(workbook, kwargs.get("sheet"), create_missing=True)
    written, formulas = _apply_cells(ws, kwargs.get("cells"))
    _save_workbook(workbook, path)
    rel = _rel(path, workspace)
    formula_note = f", including {formulas} formula(s)" if formulas else ""
    return _ok(
        f"Wrote {written} cell(s) on sheet {ws.title!r} of {rel}{formula_note}.",
        rel,
        sheet=ws.title,
        cells_written=written,
        formulas=formulas,
    )


def _action_formula(path: Path, kwargs: dict[str, Any], workspace: Path) -> ToolResult:
    payload = _maybe_json(kwargs.get("formula_cells") or kwargs.get("cells"))
    if not isinstance(payload, dict) or not payload:
        raise ValueError('`formula_cells` is required for formula (e.g. {"D2": "=SUM(B2:C2)"})')
    normalised: dict[str, str] = {}
    for key, value in payload.items():
        text = str(value).strip()
        if not text.startswith("="):
            text = "=" + text
        normalised[key] = text
    workbook = _open_workbook(path)
    ws = _pick_sheet(workbook, kwargs.get("sheet"), create_missing=True)
    written, formulas = _write_mapping(ws, normalised)
    _save_workbook(workbook, path)
    rel = _rel(path, workspace)
    return _ok(
        f"Wrote {formulas} formula(s) ({written} cell(s)) on sheet {ws.title!r} of {rel}.",
        rel,
        sheet=ws.title,
        formulas=formulas,
    )


def _action_style(path: Path, kwargs: dict[str, Any], workspace: Path) -> ToolResult:
    _Workbook, _load, Font, PatternFill, _letter, range_boundaries = _load_openpyxl()
    specs = _maybe_json(kwargs.get("styles"))
    if not isinstance(specs, list) or not specs:
        raise ValueError("`styles` must be a non-empty list of style objects")
    workbook = _open_workbook(path)
    ws = _pick_sheet(workbook, kwargs.get("sheet"), create_missing=True)
    applied = 0
    for spec in specs:
        if not isinstance(spec, dict):
            raise ValueError("each style entry must be an object")
        target = str(spec.get("target") or "").strip()
        if not target:
            raise ValueError("style.target is required (e.g. A1 or A1:D1)")
        min_col, min_row, max_col, max_row = range_boundaries(target)
        font_kwargs: dict[str, Any] = {}
        if "bold" in spec:
            font_kwargs["bold"] = _as_bool(spec.get("bold"))
        if spec.get("color"):
            font_kwargs["color"] = _rgb(str(spec["color"]))
        if spec.get("font_size") is not None:
            try:
                font_kwargs["size"] = float(spec["font_size"])
            except (TypeError, ValueError) as exc:
                raise ValueError(f"invalid font_size {spec.get('font_size')!r}") from exc
        fill = None
        if spec.get("bg"):
            fill = PatternFill(fill_type="solid", fgColor=_rgb(str(spec["bg"])))
        font = Font(**font_kwargs) if font_kwargs else None
        for row in ws.iter_rows(min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col):
            for cell in row:
                if font is not None:
                    cell.font = font
                if fill is not None:
                    cell.fill = fill
                applied += 1
    _save_workbook(workbook, path)
    rel = _rel(path, workspace)
    return _ok(
        f"Applied {len(specs)} style rule(s) to {applied} cell(s) on {ws.title!r} of {rel}.",
        rel,
        sheet=ws.title,
        styles=len(specs),
        cells_styled=applied,
    )


def _action_chart(path: Path, kwargs: dict[str, Any], workspace: Path) -> ToolResult:
    try:
        from openpyxl.chart import BarChart, LineChart, PieChart, Reference
    except ImportError as exc:  # pragma: no cover
        raise ValueError("openpyxl.chart is unavailable") from exc
    _Workbook, _load, _Font, _Fill, get_column_letter, range_boundaries = _load_openpyxl()
    spec = _maybe_json(kwargs.get("chart"))
    if not isinstance(spec, dict):
        raise ValueError(
            '`chart` is required: {type:"bar|line|pie", data_range:"A1:B5", title:"..."}'
        )
    chart_type = str(spec.get("type") or "bar").strip().lower()
    mapping = {"bar": BarChart, "line": LineChart, "pie": PieChart}
    if chart_type not in mapping:
        raise ValueError("chart.type must be one of: bar, line, pie")
    data_range = str(spec.get("data_range") or "").strip()
    if not data_range:
        raise ValueError("chart.data_range is required (e.g. A1:B5)")
    min_col, min_row, max_col, max_row = range_boundaries(data_range)
    workbook = _open_workbook(path)
    ws = _pick_sheet(workbook, kwargs.get("sheet"), create_missing=True)
    chart = mapping[chart_type]()
    title = str(spec.get("title") or "").strip()
    if title:
        chart.title = title
    data = Reference(ws, min_col=min_col, min_row=min_row, max_col=max_col, max_row=max_row)
    chart.add_data(data, titles_from_data=True)
    if chart_type == "pie" and max_col > min_col:
        cats = Reference(ws, min_col=min_col, min_row=min_row + 1, max_row=max_row)
        chart.set_categories(cats)
    anchor = f"{get_column_letter(max_col + 2)}{min_row}"
    ws.add_chart(chart, anchor)
    _save_workbook(workbook, path)
    rel = _rel(path, workspace)
    return _ok(
        f"Added {chart_type} chart on sheet {ws.title!r} of {rel} (data {data_range}).",
        rel,
        sheet=ws.title,
        chart_type=chart_type,
        data_range=data_range,
    )


def _sheet_csv(ws: Any, read_range: str | None) -> str:
    _Workbook, _load, _Font, _Fill, _letter, range_boundaries = _load_openpyxl()
    if read_range:
        min_col, min_row, max_col, max_row = range_boundaries(read_range)
    else:
        min_col, min_row = 1, 1
        max_row = ws.max_row or 1
        max_col = ws.max_column or 1
    buf = StringIO()
    writer = csv.writer(buf, lineterminator="\n")
    for row in ws.iter_rows(
        min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col, values_only=False
    ):
        writer.writerow("" if cell.value is None else cell.value for cell in row)
    return buf.getvalue()


def _action_read(path: Path, kwargs: dict[str, Any], workspace: Path) -> ToolResult:
    workbook = _open_workbook(path)
    ws = _pick_sheet(workbook, kwargs.get("sheet"))
    read_range = str(kwargs.get("read_range") or "").strip() or None
    csv_text = _sheet_csv(ws, read_range)
    rel = _rel(path, workspace)
    span = read_range or "used range"
    body = csv_text if csv_text.strip() else "(empty range)"
    return _ok(
        f"Read {rel} sheet {ws.title!r} ({span}):\n{body}",
        rel,
        sheet=ws.title,
        read_range=span,
    )


def _split_table_row(line: str) -> list[str]:
    return [part.strip() for part in line.strip().strip("|").split("|")]


def _iter_markdown_blocks(content: str) -> list[tuple[str, Any]]:
    blocks: list[tuple[str, Any]] = []
    lines = content.replace("\r\n", "\n").split("\n")
    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()
        if not stripped:
            i += 1
            continue
        heading = _HEADING.match(stripped)
        if heading:
            blocks.append(("heading", (len(heading.group(1)), heading.group(2).strip())))
            i += 1
            continue
        ul = _UL_ITEM.match(stripped)
        if ul:
            items = [ul.group(1).strip()]
            i += 1
            while i < len(lines):
                nxt = _UL_ITEM.match(lines[i].strip())
                if not nxt:
                    break
                items.append(nxt.group(1).strip())
                i += 1
            blocks.append(("list", items))
            continue
        if stripped.startswith("|"):
            rows: list[list[str]] = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                raw = lines[i].strip()
                if not _TABLE_SEP.match(raw):
                    rows.append(_split_table_row(raw))
                i += 1
            if rows:
                blocks.append(("table", rows))
            continue
        para: list[str] = [stripped]
        i += 1
        while (
            i < len(lines)
            and lines[i].strip()
            and not lines[i].strip().startswith(("#", "-", "*", "|"))
        ):
            para.append(lines[i].strip())
            i += 1
        blocks.append(("paragraph", " ".join(para)))
    return blocks


def _action_export_doc(
    path: Path, kwargs: dict[str, Any], workspace: Path, cwd_fallback: bool
) -> ToolResult:
    try:
        from docx import Document
    except ImportError as exc:  # pragma: no cover
        raise ValueError(
            "python-docx is not installed; install knorvia with the documents extra."
        ) from exc
    content = str(kwargs.get("content") or "")
    if not content.strip():
        raise ValueError("`content` is required for export_doc (Markdown subset)")
    path = _ensure_suffix(path, ".docx")
    document = Document()
    for kind, payload in _iter_markdown_blocks(content):
        if kind == "heading":
            level, text = payload
            document.add_heading(text, level=level)
        elif kind == "list":
            for item in payload:
                document.add_paragraph(item, style="List Bullet")
        elif kind == "table":
            rows: list[list[str]] = payload
            cols = max(len(row) for row in rows)
            table = document.add_table(rows=len(rows), cols=cols)
            for r_idx, row in enumerate(rows):
                for c_idx in range(cols):
                    table.cell(r_idx, c_idx).text = row[c_idx] if c_idx < len(row) else ""
        else:
            document.add_paragraph(str(payload))
    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(path)
    rel = _rel(path, workspace)
    extra = f" (absolute {path})" if cwd_fallback else ""
    return _ok(
        f"Exported Word document {rel} ({path.stat().st_size} bytes){extra}.",
        rel,
        bytes=path.stat().st_size,
    )


def _iter_slides(content: str) -> list[tuple[str, list[str]]]:
    """Parse an outline into (title, bullets) slides.

    Markdown headings (# ...) start a new slide — the documented format.
    Fallback for plain outlines with NO headings at all: every non-bullet
    top-level line starts a slide and collects following "- " lines as its
    bullets. Mixed input keeps heading semantics untouched.
    """
    lines = content.replace("\r\n", "\n").split("\n")
    has_headings = any(_HEADING.match(line.strip()) for line in lines)
    if not has_headings:
        return _iter_plain_outline(lines)

    slides: list[tuple[str, list[str]]] = []
    title: str | None = None
    bullets: list[str] = []

    def flush() -> None:
        nonlocal title, bullets
        if title is None and not bullets:
            return
        slides.append((title or "Slide", bullets))
        title = None
        bullets = []

    for raw in lines:
        stripped = raw.strip()
        if not stripped:
            continue
        heading = _HEADING.match(stripped)
        if heading:
            flush()
            title = heading.group(2).strip()
            bullets = []
            continue
        if title is None:
            title = stripped.lstrip("# ").strip()
            continue
        bullets.append(stripped)
    flush()
    return slides


def _iter_plain_outline(lines: list[str]) -> list[tuple[str, list[str]]]:
    """No-markdown fallback: top-level line = slide title; '- ' lines = bullets."""
    slides: list[tuple[str, list[str]]] = []
    title: str | None = None
    bullets: list[str] = []

    def flush() -> None:
        nonlocal title, bullets
        if title is None:
            return
        slides.append((title, bullets))
        title = None
        bullets = []

    for raw in lines:
        stripped = raw.strip()
        if not stripped:
            continue
        if stripped.startswith("- ") or stripped.startswith("•"):
            if title is not None:
                bullets.append(stripped.lstrip("-• ").strip())
            continue
        flush()
        title = stripped
    flush()
    return slides or [("Slide", [])]


def _action_export_slide(
    path: Path, kwargs: dict[str, Any], workspace: Path, cwd_fallback: bool
) -> ToolResult:
    try:
        from pptx import Presentation
        from pptx.util import Inches, Pt
    except ImportError as exc:  # pragma: no cover
        raise ValueError(
            "python-pptx is not installed; install knorvia with the documents extra."
        ) from exc
    content = str(kwargs.get("content") or "")
    if not content.strip():
        raise ValueError("`content` is required for export_slide (outline text)")
    slides = _iter_slides(content)
    if not slides:
        raise ValueError("export_slide produced no slides from `content`")
    path = _ensure_suffix(path, ".pptx")
    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)
    blank = presentation.slide_layouts[6]
    for title, bullets in slides:
        slide = presentation.slides.add_slide(blank)
        title_box = slide.shapes.add_textbox(Inches(0.7), Inches(0.4), Inches(11.9), Inches(1.1))
        title_tf = title_box.text_frame
        title_tf.text = title
        if title_tf.paragraphs:
            title_tf.paragraphs[0].font.size = Pt(32)
            title_tf.paragraphs[0].font.bold = True
        body = slide.shapes.add_textbox(Inches(0.7), Inches(1.7), Inches(11.9), Inches(5.2))
        body_tf = body.text_frame
        body_tf.word_wrap = True
        if not bullets:
            body_tf.text = ""
        else:
            for idx, bullet in enumerate(bullets):
                paragraph = body_tf.paragraphs[0] if idx == 0 else body_tf.add_paragraph()
                paragraph.text = bullet
                paragraph.level = 0
                paragraph.font.size = Pt(20)
    path.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(path)
    rel = _rel(path, workspace)
    extra = f" (absolute {path})" if cwd_fallback else ""
    return _ok(
        f"Exported {len(slides)}-slide deck {rel} (16:9){extra}.",
        rel,
        slides=len(slides),
        bytes=path.stat().st_size,
    )


def _render_preview_png(ws: Any, dest: Path, *, max_rows: int = 16, max_cols: int = 8) -> bool:
    try:
        from PIL import Image, ImageDraw, ImageFont
    except ImportError:
        return False
    rows = max(1, min(int(ws.max_row or 1), max_rows))
    cols = max(1, min(int(ws.max_column or 1), max_cols))
    cell_w, cell_h, pad = 96, 24, 8
    image = Image.new("RGB", (cols * cell_w + pad * 2, rows * cell_h + pad * 2), "white")
    draw = ImageDraw.Draw(image)
    try:
        font = ImageFont.truetype("arial.ttf", 12)
    except OSError:
        font = ImageFont.load_default()
    for row in range(1, rows + 1):
        for col in range(1, cols + 1):
            x0 = pad + (col - 1) * cell_w
            y0 = pad + (row - 1) * cell_h
            fill = "#F3EDE6" if row == 1 else "white"
            draw.rectangle([x0, y0, x0 + cell_w, y0 + cell_h], outline="#C8C8C8", fill=fill)
            value = ws.cell(row, col).value
            text = "" if value is None else str(value)[:14]
            draw.text((x0 + 4, y0 + 5), text, fill="#222222", font=font)
    dest.parent.mkdir(parents=True, exist_ok=True)
    image.save(dest, format="PNG")
    return dest.is_file() and dest.stat().st_size > 0


def _action_screenshot_hint(path: Path, kwargs: dict[str, Any], workspace: Path) -> ToolResult:
    workbook = _open_workbook(path)
    ws = _pick_sheet(workbook, kwargs.get("sheet"))
    read_range = str(kwargs.get("read_range") or "").strip() or None
    grid = _sheet_csv(ws, read_range)
    preview = path.with_name(f"{path.stem}_preview.png")
    png_ok = _render_preview_png(ws, preview)
    rel = _rel(path, workspace)
    preview_rel = _rel(preview, workspace) if png_ok else ""
    output = preview_rel or rel
    png_note = f" PNG preview: {preview_rel}." if png_ok else " (PNG preview unavailable)."
    body = grid if grid.strip() else "(empty sheet)"
    return _ok(
        f"Screenshot hint for {rel} sheet {ws.title!r}:{png_note}\n{body}",
        output,
        sheet=ws.title,
        preview_file=preview_rel,
        source_file=rel,
    )


def _task_dir_of(kwargs: dict[str, Any]) -> Path | None:
    raw = str(kwargs.get("_task_dir") or "").strip()
    if not raw:
        return None
    return Path(raw).expanduser().resolve()


def _prepare_draft(
    kwargs: dict[str, Any],
    workspace: Path,
    *,
    writing: bool,
    lifecycle: str,
) -> tuple[OfficeDraftStore | None, str, Path]:
    """Resolve the draft store and the directory write actions should use.

    ``as_draft`` defaults to true, but degrades to a direct workspace write
    when the pipeline did not inject a turn ``_task_dir`` (unit tests / CWD
    fallback). Terminal drafts refuse further writes.
    """
    as_draft = _as_bool(kwargs.get("as_draft"), default=True)
    task_dir = _task_dir_of(kwargs)
    draft_id = str(kwargs.get("draft_id") or kwargs.get("_office_draft_id") or "").strip()
    needs_store = bool(lifecycle) or (writing and as_draft)
    if not needs_store:
        return None, "", workspace
    if task_dir is None:
        if writing:
            return None, "", workspace
        raise ValueError("office drafts need a turn workspace (`_task_dir`).")
    store = OfficeDraftStore(task_dir, workspace_dir=workspace)
    if lifecycle == "create" and not draft_id:
        draft_id = store.create()
    if writing and as_draft:
        if not draft_id:
            draft_id = store.create()
        else:
            store.assert_writable(draft_id)
        return store, draft_id, store.draft_dir(draft_id)
    if lifecycle in LIFECYCLE_ACTIONS:
        if not draft_id:
            raise ValueError("`draft_id` is required for this draft action.")
        return store, draft_id, workspace
    return store, draft_id, workspace


def _resolve_source_file(
    write_ws: Path,
    workspace: Path,
    source_name: str,
) -> Path:
    """Locate ``source_file`` in the draft write dir first, then the workspace."""
    raw = str(source_name or "").strip()
    if not raw:
        raise ValueError("`source_file` is required when provided")
    for root in (write_ws, workspace):
        try:
            candidate = _resolve_file(root, raw)
        except ValueError:
            continue
        if candidate.is_file():
            return candidate
    raise ValueError(f"source_file not found: {raw}")


def _parse_refs(raw: Any) -> list[dict[str, Any]]:
    payload = _maybe_json(raw)
    if payload in (None, "", [], {}):
        return []
    if isinstance(payload, dict):
        payload = [payload]
    if not isinstance(payload, list):
        raise ValueError('`refs` must be a list like [{"to":"sheet","range":"A1:B5"}]')
    refs: list[dict[str, Any]] = []
    for item in payload:
        if not isinstance(item, dict):
            raise ValueError("each refs entry must be an object")
        refs.append(item)
    return refs


def _action_container_new(
    path: Path,
    kwargs: dict[str, Any],
    workspace: Path,
    cwd_fallback: bool,
    *,
    write_ws: Path | None = None,
    official_ws: Path | None = None,
) -> ToolResult:
    path = _ensure_suffix(path, ".univer")
    units_raw = _maybe_json(kwargs.get("units") or kwargs.get("cells"))
    entries: list[dict[str, Any]] = []
    if isinstance(units_raw, list):
        for item in units_raw:
            if not isinstance(item, dict):
                raise ValueError("each units entry must be an object")
            entries.append(dict(item))
    refs = _parse_refs(kwargs.get("refs"))
    # Optional seed: pack an existing native file as the first unit.
    source_name = str(kwargs.get("source_file") or "").strip()
    if source_name:
        source_path = _resolve_source_file(
            write_ws or workspace,
            official_ws or workspace,
            source_name,
        )
        suffix = source_path.suffix.lower()
        type_map = {".xlsx": "sheet", ".xlsm": "sheet", ".docx": "doc", ".pptx": "slide"}
        unit_type = type_map.get(suffix)
        if unit_type is None:
            raise ValueError("source_file for container_new must be .xlsx / .docx / .pptx")
        unit_id = str(kwargs.get("unit_id") or "").strip()
        if not unit_id:
            stem = Path(source_name).stem
            unit_id = stem if re.fullmatch(r"[A-Za-z][A-Za-z0-9_-]{0,31}", stem) else unit_type
        entries.append(
            {
                "id": unit_id,
                "type": unit_type,
                "name": str(kwargs.get("unit_name") or kwargs.get("sheet") or unit_id),
                "source": source_path,
            }
        )
    manifest = pack_univer(entries, path, refs=refs or None)
    rel = _rel(path, workspace)
    extra = f" (absolute {path})" if cwd_fallback else ""
    count = len(manifest.get("units") or [])
    return _ok(
        f"Created .univer container {rel} with {count} unit(s){extra}.",
        rel,
        units=manifest.get("units") or [],
        refs=manifest.get("refs") or [],
    )


def _action_container_add(
    path: Path,
    kwargs: dict[str, Any],
    workspace: Path,
    write_ws: Path,
    official_ws: Path,
) -> ToolResult:
    path = _ensure_suffix(path, ".univer")
    if not path.is_file():
        raise ValueError(
            f"Container {path.name!r} does not exist. Call action=container_new first."
        )
    unit_type = str(kwargs.get("unit_type") or "").strip().lower()
    if unit_type not in UNIT_TYPES:
        raise ValueError(f"`unit_type` is required for container_add ({', '.join(UNIT_TYPES)})")
    unit_name = (
        str(kwargs.get("unit_name") or kwargs.get("name") or kwargs.get("sheet") or "").strip()
        or None
    )
    unit_id = str(kwargs.get("unit_id") or "").strip() or None
    source_bytes: bytes | None = None
    source_name = str(kwargs.get("source_file") or "").strip()
    if source_name:
        source_path = _resolve_source_file(write_ws, official_ws, source_name)
        source_bytes = source_path.read_bytes()
    refs = _parse_refs(kwargs.get("refs"))
    unit = add_unit(
        path,
        unit_type,
        unit_name,
        source_bytes,
        unit_id=unit_id,
        refs=refs or None,
    )
    rel = _rel(path, workspace)
    ref_note = ""
    if refs:
        ref_note = f", recorded {len(refs)} ref(s)"
    return _ok(
        f"Added {unit['type']} unit {unit['id']!r} to {rel}{ref_note}.",
        rel,
        unit=unit,
        units=list_units(path),
        refs=open_univer(path).get("refs") or [],
    )


def _action_container_export_unit(
    path: Path,
    kwargs: dict[str, Any],
    workspace: Path,
    write_ws: Path,
) -> ToolResult:
    path = _ensure_suffix(path, ".univer")
    if not path.is_file():
        raise ValueError(f"Container {path.name!r} does not exist.")
    unit_id = str(kwargs.get("unit_id") or "").strip()
    if not unit_id:
        raise ValueError("`unit_id` is required for container_export_unit")
    out_name = str(kwargs.get("export_file") or kwargs.get("output") or "").strip()
    if not out_name:
        out_name = unit_id
    out_path = _resolve_file(write_ws, out_name)
    exported = export_unit(path, unit_id, out_path)
    rel = _rel(exported, write_ws)
    return _ok(
        f"Exported unit {unit_id!r} from {path.name} to {rel}.",
        rel,
        unit_id=unit_id,
        container=_rel(path, workspace),
    )


def _run_write_action(
    action: str,
    path: Path,
    kwargs: dict[str, Any],
    workspace: Path,
    cwd_fallback: bool,
    *,
    write_ws: Path | None = None,
    official_ws: Path | None = None,
) -> ToolResult:
    active_ws = write_ws or workspace
    root_ws = official_ws or workspace
    if action == "create":
        return _action_create(path, kwargs, workspace, cwd_fallback)
    if action == "add_sheet":
        return _action_add_sheet(_ensure_suffix(path, ".xlsx"), kwargs, workspace)
    if action == "write_cells":
        return _action_write_cells(_ensure_suffix(path, ".xlsx"), kwargs, workspace)
    if action == "formula":
        return _action_formula(_ensure_suffix(path, ".xlsx"), kwargs, workspace)
    if action == "style":
        return _action_style(_ensure_suffix(path, ".xlsx"), kwargs, workspace)
    if action == "chart":
        return _action_chart(_ensure_suffix(path, ".xlsx"), kwargs, workspace)
    if action == "read":
        return _action_read(_ensure_suffix(path, ".xlsx"), kwargs, workspace)
    if action == "export_doc":
        return _action_export_doc(path, kwargs, workspace, cwd_fallback)
    if action == "export_slide":
        return _action_export_slide(path, kwargs, workspace, cwd_fallback)
    if action == "screenshot_hint":
        return _action_screenshot_hint(_ensure_suffix(path, ".xlsx"), kwargs, workspace)
    if action == "container_new":
        return _action_container_new(
            path,
            kwargs,
            workspace,
            cwd_fallback,
            write_ws=active_ws,
            official_ws=root_ws,
        )
    if action == "container_add":
        return _action_container_add(path, kwargs, workspace, active_ws, root_ws)
    if action == "container_export_unit":
        return _action_container_export_unit(path, kwargs, workspace, active_ws)
    return _fail(f"Unhandled action {action!r}.")


def _run_lifecycle(store: OfficeDraftStore, draft_id: str, action: str) -> dict[str, Any]:
    if action == "create":
        return store.status(draft_id)
    if action == "ready":
        return store.mark_ready(draft_id)
    if action == "merge":
        return store.merge(draft_id)
    if action == "discard":
        return store.discard(draft_id)
    if action == "status":
        meta = store.status(draft_id)
        meta["diff"] = store.diff(draft_id)
        return meta
    raise ValueError(f"Unhandled draft action {action!r}.")


def _attach_draft_meta(
    result: ToolResult,
    store: OfficeDraftStore | None,
    draft_id: str,
) -> ToolResult:
    if store is None or not draft_id:
        return result
    try:
        payload = store.card_payload(draft_id)
    except DraftError:
        return result
    merged = {**(result.metadata or {}), **payload}
    return ToolResult(content=result.content, success=result.success, metadata=merged)


def execute_office_document(kwargs: dict[str, Any]) -> ToolResult:
    """Run one ``office_document`` action. Public entry used by the tool class."""
    action = str(kwargs.get("action") or "").strip().lower()
    draft_action = str(kwargs.get("draft_action") or "").strip().lower()
    if draft_action and draft_action not in DRAFT_ACTIONS:
        valid = ", ".join(DRAFT_ACTIONS)
        return _fail(f"Invalid draft_action {draft_action!r}. Valid: {valid}.")
    if action and action not in ACTIONS:
        valid = ", ".join(ACTIONS)
        return _fail(f"Invalid action {action!r}. Valid actions: {valid}.")
    if not action and not draft_action:
        valid = ", ".join(ACTIONS)
        return _fail(f"Invalid action {action!r}. Valid actions: {valid}.")
    lifecycle = (
        draft_action
        if draft_action in DRAFT_ACTIONS
        else (action if action in LIFECYCLE_ACTIONS else "")
    )
    writing = action in WRITE_ACTIONS
    workspace, cwd_fallback = _workspace(kwargs)
    try:
        store, draft_id, write_ws = _prepare_draft(
            kwargs, workspace, writing=writing, lifecycle=lifecycle
        )
        result: ToolResult | None = None
        if writing:
            file_name = str(kwargs.get("file") or "").strip()
            if not file_name:
                return _fail("`file` is required (relative to this turn's workspace).")
            path = _resolve_file(write_ws, file_name)
            result = _run_write_action(
                action,
                path,
                kwargs,
                write_ws,
                cwd_fallback,
                write_ws=write_ws,
                official_ws=workspace,
            )
            if (
                result.success
                and store is not None
                and draft_id
                and _as_bool(kwargs.get("as_draft"), default=True)
            ):
                output = str((result.metadata or {}).get("output_file") or "")
                if output:
                    store.note_file(draft_id, output)
                preview = str((result.metadata or {}).get("preview_file") or "")
                if preview:
                    store.note_file(draft_id, preview)
                if action in CONTAINER_ACTIONS and path.suffix.lower() == ".univer":
                    store.note_file(draft_id, _rel(path, write_ws))
        if lifecycle and lifecycle != "create":
            if store is None or not draft_id:
                return _fail("`draft_id` is required for this draft action.")
            meta = _run_lifecycle(store, draft_id, lifecycle)
            summary = (
                f"Office draft {draft_id} is {meta.get('status')}."
                if result is None
                else f"{result.content} Draft {draft_id} is now {meta.get('status')}."
            )
            success = True if result is None else result.success
            extra = dict(result.metadata or {}) if result is not None else {}
            result = ToolResult(content=summary, success=success, metadata=extra)
        elif lifecycle == "create" and result is None:
            if store is None or not draft_id:
                return _fail("Could not create an office draft.")
            result = ToolResult(
                content=f"Created office draft {draft_id}.",
                success=True,
                metadata={},
            )
        if result is None:
            return _fail(f"Unhandled action {action!r}.")
        return _attach_draft_meta(result, store, draft_id)
    except (ValueError, ContainerError) as exc:
        return _fail(str(exc))
    except Exception as exc:
        return _fail(f"office_document failed: {exc}")


class OfficeDocumentTool(BaseTool):
    """Structured xlsx / docx / pptx authoring for the agentic chat loop."""

    def get_prompt_hints(self, language: str = "en"):
        return load_prompt_hints(self.name, language=language)

    def get_definition(self) -> ToolDefinition:
        return ToolDefinition(
            name="office_document",
            description=_TOOL_DESCRIPTION,
            raw_parameters={
                "type": "object",
                "additionalProperties": True,
                "required": ["action"],
                "properties": {
                    "action": {
                        "type": "string",
                        "enum": list(ACTIONS),
                        "description": (
                            "create | add_sheet | write_cells | formula | style | "
                            "chart | read | export_doc | export_slide | screenshot_hint "
                            "| container_new | container_add | container_export_unit "
                            "| ready | merge | discard | status"
                        ),
                    },
                    "as_draft": {
                        "type": "boolean",
                        "description": (
                            "Write into an isolated draft the user must confirm "
                            "(default true). Set false only to write the official file."
                        ),
                    },
                    "draft_action": {
                        "type": "string",
                        "enum": list(DRAFT_ACTIONS),
                        "description": (
                            "Independent draft lifecycle: create | ready | merge | "
                            "discard | status. ready/merge/discard/status also work as `action`."
                        ),
                    },
                    "draft_id": {
                        "type": "string",
                        "description": (
                            "Existing 8-char draft id. Reused automatically within a turn "
                            "when omitted."
                        ),
                    },
                    "file": {
                        "type": "string",
                        "description": "Output file name relative to this turn's workspace. Required for write actions.",
                    },
                    "sheet": {
                        "type": "string",
                        "description": "Worksheet / slide name.",
                    },
                    "cells": {
                        "description": (
                            'Object {"A1": "Title", "B2": 123} or 2D array [["a","b"],[1,2]].'
                        ),
                    },
                    "formula_cells": {
                        "type": "object",
                        "additionalProperties": {"type": "string"},
                        "description": '{"D2": "=SUM(B2:C2)"}',
                    },
                    "styles": {
                        "type": "array",
                        "items": {"type": "object", "additionalProperties": True},
                        "description": (
                            '[{target:"A1:D1", bold:true, bg:"#B0501E", '
                            'color:"#FFFFFF", font_size:12}]'
                        ),
                    },
                    "chart": {
                        "type": "object",
                        "additionalProperties": True,
                        "description": '{type:"bar|line|pie", data_range:"A1:B5", title:"..."}',
                    },
                    "content": {
                        "type": "string",
                        "description": "Markdown (export_doc) or outline (export_slide).",
                    },
                    "read_range": {
                        "type": "string",
                        "description": 'A1-style range such as "A1:D10".',
                    },
                    "unit_type": {
                        "type": "string",
                        "enum": list(UNIT_TYPES),
                        "description": (
                            "container_add: sheet | doc | slide unit to append "
                            "to a .univer ZIP container."
                        ),
                    },
                    "unit_id": {
                        "type": "string",
                        "description": (
                            "Unit id inside a .univer container "
                            "(container_add / container_export_unit)."
                        ),
                    },
                    "unit_name": {
                        "type": "string",
                        "description": "Display name for a container unit.",
                    },
                    "source_file": {
                        "type": "string",
                        "description": (
                            "Existing .xlsx/.docx/.pptx in the turn workspace to "
                            "pack into a .univer unit."
                        ),
                    },
                    "refs": {
                        "type": "array",
                        "items": {"type": "object", "additionalProperties": True},
                        "description": (
                            'Cross-unit data refs, e.g. [{"to":"sheet","range":"A1:B5"}]. '
                            "Recorded on manifest.refs; Slide units get a text placeholder."
                        ),
                    },
                    "export_file": {
                        "type": "string",
                        "description": (
                            "Destination file name for container_export_unit "
                            "(native xlsx/docx/pptx)."
                        ),
                    },
                },
            },
        )

    async def execute(self, **kwargs: Any) -> ToolResult:
        return execute_office_document(kwargs)

"""XLSX read model over openpyxl.

openpyxl is used for *reading only*. Saving an imported workbook through
openpyxl is explicitly forbidden (see ``xlsx_patch``); generated workbooks
are written through the adapter.
"""

from __future__ import annotations

from datetime import date, datetime
from io import BytesIO
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter, range_boundaries

from knorvia.services.office_artifacts.contracts import (
    MAX_READ_CELLS,
    InvalidOperationError,
    cell_name,
    parse_range,
)


class XlsxReader:
    """Read-model facade over one in-memory workbook."""

    def __init__(self, data: bytes) -> None:
        try:
            self._formulas = load_workbook(BytesIO(data), data_only=False)
            self._cached = load_workbook(BytesIO(data), data_only=True)
        except Exception as exc:  # noqa: BLE001 - openpyxl raises bare exceptions
            raise InvalidOperationError(f"workbook could not be opened for reading: {exc}")

    # ------------------------------------------------------------------
    # Sheet helpers
    # ------------------------------------------------------------------

    def sheet_names(self) -> list[str]:
        return list(self._formulas.sheetnames)

    def _worksheet(self, sheet: str | None, cached: bool = False):
        wb = self._cached if cached else self._formulas
        if sheet:
            if sheet not in wb.sheetnames:
                raise InvalidOperationError(
                    f"sheet {sheet!r} not found; existing: {', '.join(wb.sheetnames)}"
                )
            return wb[sheet]
        return wb.active

    # ------------------------------------------------------------------
    # Overview / features
    # ------------------------------------------------------------------

    def overview(self) -> dict[str, Any]:
        sheets = []
        for name in self._formulas.sheetnames:
            ws = self._formulas[name]
            sheets.append(
                {
                    "name": name,
                    "state": str(ws.sheet_state or "visible"),
                    "used_range": self._used_range(ws),
                    "merged_ranges": [str(rng) for rng in ws.merged_cells.ranges][:50],
                    "freeze_panes": str(ws.freeze_panes) if ws.freeze_panes else None,
                }
            )
        return {
            "sheet_names": self._formulas.sheetnames,
            "sheets": sheets,
            "defined_names": sorted(str(n) for n in getattr(self._formulas, "defined_names", {})),
        }

    def features(self, sheet: str | None = None) -> dict[str, Any]:
        names = [sheet] if sheet else self._formulas.sheetnames
        report: dict[str, Any] = {}
        for name in names:
            ws = self._worksheet(name)
            report[name] = {
                "tables": sorted(ws.tables.keys()) if hasattr(ws, "tables") else [],
                "charts": len(getattr(ws, "_charts", []) or []),
                "images": len(getattr(ws, "_images", []) or []),
                "data_validations": len(getattr(ws, "data_validations", None)
                                        and ws.data_validations.dataValidation or []),
                "conditional_formatting": len(list(getattr(ws, "conditional_formatting", []) or [])),
                "named_ranges": len(
                    [n for n in getattr(ws.parent, "defined_names", {}).values()
                     if getattr(n, "destination", None)]
                ),
                "merged_count": len(ws.merged_cells.ranges),
                "hidden": ws.sheet_state != "visible",
                "protected": ws.protection.sheet is True
                if hasattr(ws, "protection")
                else False,
                "freeze_panes": str(ws.freeze_panes) if ws.freeze_panes else None,
                "hyperlinks": len([c for row in ws.iter_rows() for c in row if c.hyperlink]),
                "comments": len([c for row in ws.iter_rows() for c in row if c.comment]),
            }
        return report

    @staticmethod
    def _used_range(ws: Any) -> str:
        max_row = ws.max_row or 1
        max_col = ws.max_column or 1
        return f"A1:{get_column_letter(max_col)}{max_row}"

    # ------------------------------------------------------------------
    # Range reads
    # ------------------------------------------------------------------

    def read_range(self, sheet: str | None, range_str: str | None) -> dict[str, Any]:
        ws = self._worksheet(sheet)
        ws_cached = self._worksheet(ws.title, cached=True)
        if range_str:
            min_col, min_row, max_col, max_row = parse_range(range_str)
        else:
            min_col, min_row = 1, 1
            max_row, max_col = ws.max_row or 1, ws.max_column or 1
        total = (max_col - min_col + 1) * (max_row - min_row + 1)
        if total > MAX_READ_CELLS:
            raise InvalidOperationError(
                f"range covers {total} cells; reads are limited to {MAX_READ_CELLS}"
            )
        merged_lookup = {
            str(rng): rng for rng in ws.merged_cells.ranges
        }
        cells: list[dict[str, Any]] = []
        for row in range(min_row, max_row + 1):
            for col in range(min_col, max_col + 1):
                address = cell_name(col, row)
                cell = ws.cell(row=row, column=col)
                cached_cell = ws_cached.cell(row=row, column=col)
                merge_range = next(
                    (str(r) for r, rng in merged_lookup.items() if address in rng), None
                )
                value = cell.value
                is_formula = isinstance(value, str) and value.startswith("=")
                cells.append(
                    {
                        "cell": address,
                        "value": None if is_formula else _jsonable(value),
                        "formula": value if is_formula else "",
                        "cached_value": _jsonable(cached_cell.value) if is_formula else None,
                        "style_id": str(getattr(cell, "style_id", "") or ""),
                        "merge_range": merge_range,
                    }
                )
        return {
            "sheet": ws.title,
            "range": f"{cell_name(min_col, min_row)}:{cell_name(max_col, max_row)}",
            "cells": cells,
            "count": len(cells),
        }

    def snapshot_cells(self, sheet: str | None, targets: list[str]) -> dict[str, dict[str, Any]]:
        """{A1: snapshot} for the given targets, used for semantic diffs."""
        ws = self._worksheet(sheet)
        ws_cached = self._worksheet(ws.title, cached=True)
        snapshots: dict[str, dict[str, Any]] = {}
        for target in targets:
            col, row = _parse_cell_strict(target)
            cell = ws.cell(row=row, column=col)
            cached_cell = ws_cached.cell(row=row, column=col)
            value = cell.value
            is_formula = isinstance(value, str) and value.startswith("=")
            snapshots[target] = {
                "value": None if is_formula else _jsonable(value),
                "formula": value if is_formula else "",
                "cached_value": _jsonable(cached_cell.value) if is_formula else None,
                "style_id": str(getattr(cell, "style_id", "") or ""),
            }
        return snapshots

    def find(
        self,
        query: str,
        *,
        sheet: str | None = None,
        scope: str | None = None,
        max_hits: int = 50,
    ) -> list[dict[str, Any]]:
        needle = str(query or "").lower()
        if not needle:
            raise InvalidOperationError("find requires a non-empty query")
        names = [sheet] if sheet else self._formulas.sheetnames
        hits: list[dict[str, Any]] = []
        for name in names:
            ws = self._worksheet(name)
            if scope:
                min_col, min_row, max_col, max_row = parse_range(scope)
            else:
                min_col, min_row, max_col, max_row = 1, 1, ws.max_column or 1, ws.max_row or 1
            for row in range(min_row, min(max_row, (ws.max_row or 1)) + 1):
                for col in range(min_col, min(max_col, (ws.max_column or 1)) + 1):
                    value = ws.cell(row=row, column=col).value
                    if value is not None and needle in str(value).lower():
                        hits.append({"sheet": name, "cell": cell_name(col, row), "value": _jsonable(value)})
                        if len(hits) >= max_hits:
                            return hits
        return hits

    def external_link_count(self) -> int:
        """External references force fail-closed for formula patching."""
        try:
            return len(self._formulas._external_links or [])
        except Exception:
            return 0


def _parse_cell_strict(target: str) -> tuple[int, int]:
    from knorvia.services.office_artifacts.contracts import parse_cell

    return parse_cell(target)


def _jsonable(value: Any) -> Any:
    if value is None or isinstance(value, (str, int, bool, float)):
        return value
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    return str(value)


__all__ = ["XlsxReader", "range_boundaries"]

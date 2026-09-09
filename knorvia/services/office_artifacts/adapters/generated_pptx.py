"""Generated-PPTX adapter: outline → 16:9 deck (from-zero creation only).

Imported PPTX files are inspect/preview/copy-to-draft in this phase; any
mutation on them fails closed. Rendering logic is moved verbatim from the
legacy ``office_document`` tool so output stays compatible.
"""

from __future__ import annotations

from pathlib import Path
import re

from knorvia.services.office_artifacts.contracts import InvalidOperationError

_HEADING = re.compile(r"^(#{1,3})\s+(.*\S)\s*$")


def _iter_slides(content: str) -> list[tuple[str, list[str]]]:
    lines = content.replace("\r\n", "\n").split("\n")
    has_headings = any(_HEADING.match(line.strip()) for line in lines)
    if not has_headings:
        return _iter_plain_outline(lines)
    slides: list[tuple[str, list[str]]] = []
    title: str | None = None
    bullets: list[str] = []

    def flush() -> None:
        nonlocal title, bullets
        if title is None and not bullets:
            return
        slides.append((title or "Slide", bullets))
        title = None
        bullets = []

    for raw in lines:
        stripped = raw.strip()
        if not stripped:
            continue
        heading = _HEADING.match(stripped)
        if heading:
            flush()
            title = heading.group(2).strip()
            bullets = []
            continue
        if title is None:
            title = stripped.lstrip("# ").strip()
            continue
        bullets.append(stripped)
    flush()
    return slides


def _iter_plain_outline(lines: list[str]) -> list[tuple[str, list[str]]]:
    slides: list[tuple[str, list[str]]] = []
    title: str | None = None
    bullets: list[str] = []

    def flush() -> None:
        nonlocal title, bullets
        if title is None:
            return
        slides.append((title, bullets))
        title = None
        bullets = []

    for raw in lines:
        stripped = raw.strip()
        if not stripped:
            continue
        if stripped.startswith("- ") or stripped.startswith("•"):
            if title is not None:
                bullets.append(stripped.lstrip("-• ").strip())
            continue
        flush()
        title = stripped
    flush()
    return slides or [("Slide", [])]


def generate_pptx(content: str) -> tuple[bytes, int]:
    try:
        from pptx import Presentation
        from pptx.util import Inches, Pt
    except ImportError as exc:  # pragma: no cover
        raise InvalidOperationError(
            "python-pptx is not installed; install knorvia with the documents extra."
        ) from exc
    if not str(content or "").strip():
        raise InvalidOperationError("`content` is required for export_slide (outline text)")
    slides = _iter_slides(content)
    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)
    blank = presentation.slide_layouts[6]
    for title, bullets in slides:
        slide = presentation.slides.add_slide(blank)
        title_box = slide.shapes.add_textbox(Inches(0.7), Inches(0.4), Inches(11.9), Inches(1.1))
        title_tf = title_box.text_frame
        title_tf.text = title
        if title_tf.paragraphs:
            title_tf.paragraphs[0].font.size = Pt(32)
            title_tf.paragraphs[0].font.bold = True
        body = slide.shapes.add_textbox(Inches(0.7), Inches(1.7), Inches(11.9), Inches(5.2))
        body_tf = body.text_frame
        body_tf.word_wrap = True
        if not bullets:
            body_tf.text = ""
        else:
            for idx, bullet in enumerate(bullets):
                paragraph = body_tf.paragraphs[0] if idx == 0 else body_tf.add_paragraph()
                paragraph.text = bullet
                paragraph.level = 0
                paragraph.font.size = Pt(20)
    buffer = _BytesIO()
    presentation.save(buffer)
    return buffer.getvalue(), len(slides)


def _BytesIO():  # noqa: N802
    import io

    return io.BytesIO()


def mutate_imported_pptx_unsupported() -> None:
    raise UnsupportedPptxMutation()


class UnsupportedPptxMutation(InvalidOperationError):
    def __init__(self) -> None:
        super().__init__(
            "editing imported .pptx files is not supported in this phase; "
            "layout-safe PPTX mutation is on the roadmap. Use inspect/preview, "
            "or create a new deck."
        )


def ensure_pptx_suffix(name: str) -> str:
    path = Path(name)
    return name if path.suffix.lower() == ".pptx" else str(path.with_suffix(".pptx"))

"""Generated-DOCX adapter: Markdown subset → .docx (from-zero creation only).

Imported DOCX files are inspect/preview/copy-to-draft in this phase; any
mutation on them fails closed. The rendering logic is moved verbatim from
the legacy ``office_document`` tool so output stays compatible.
"""

from __future__ import annotations

from pathlib import Path
import re
from typing import Any

from knorvia.services.office_artifacts.contracts import InvalidOperationError

_HEADING = re.compile(r"^(#{1,3})\s+(.*\S)\s*$")
_UL_ITEM = re.compile(r"^[-*]\s+(.*)$")
_TABLE_SEP = re.compile(r"^\|?\s*:?-+:?\s*(\|\s*:?-+:?\s*)+\|?\s*$")


def _iter_markdown_blocks(content: str) -> list[tuple[str, Any]]:
    blocks: list[tuple[str, Any]] = []
    lines = content.replace("\r\n", "\n").split("\n")
    i = 0
    while i < len(lines):
        stripped = lines[i].strip()
        if not stripped:
            i += 1
            continue
        heading = _HEADING.match(stripped)
        if heading:
            blocks.append(("heading", (len(heading.group(1)), heading.group(2).strip())))
            i += 1
            continue
        ul = _UL_ITEM.match(stripped)
        if ul:
            items = [ul.group(1).strip()]
            i += 1
            while i < len(lines) and _UL_ITEM.match(lines[i].strip()):
                items.append(_UL_ITEM.match(lines[i].strip()).group(1).strip())
                i += 1
            blocks.append(("list", items))
            continue
        if stripped.startswith("|"):
            rows: list[list[str]] = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                raw = lines[i].strip()
                if not _TABLE_SEP.match(raw):
                    rows.append([p.strip() for p in raw.strip().strip("|").split("|")])
                i += 1
            if rows:
                blocks.append(("table", rows))
            continue
        para: list[str] = [stripped]
        i += 1
        while (
            i < len(lines)
            and lines[i].strip()
            and not lines[i].strip().startswith(("#", "-", "*", "|"))
        ):
            para.append(lines[i].strip())
            i += 1
        blocks.append(("paragraph", " ".join(para)))
    return blocks


def generate_docx(content: str) -> bytes:
    try:
        from docx import Document
    except ImportError as exc:  # pragma: no cover
        raise InvalidOperationError(
            "python-docx is not installed; install knorvia with the documents extra."
        ) from exc
    if not str(content or "").strip():
        raise InvalidOperationError("`content` is required for export_doc (Markdown subset)")
    document = Document()
    for kind, payload in _iter_markdown_blocks(content):
        if kind == "heading":
            level, text = payload
            document.add_heading(text, level=level)
        elif kind == "list":
            for item in payload:
                document.add_paragraph(item, style="List Bullet")
        elif kind == "table":
            rows: list[list[str]] = payload
            cols = max(len(row) for row in rows)
            table = document.add_table(rows=len(rows), cols=cols)
            for r_idx, row in enumerate(rows):
                for c_idx in range(cols):
                    table.cell(r_idx, c_idx).text = row[c_idx] if c_idx < len(row) else ""
        else:
            document.add_paragraph(str(payload))
    buffer = BytesIO_()
    document.save(buffer)
    return buffer.getvalue()


def BytesIO_():  # noqa: N802
    import io

    return io.BytesIO()


def mutate_imported_docx_unsupported() -> None:
    raise UnsupportedDocxMutation()


class UnsupportedDocxMutation(InvalidOperationError):
    def __init__(self) -> None:
        super().__init__(
            "editing imported .docx files is not supported in this phase; "
            "the narrow-patch DOCX pipeline is on the roadmap. Use inspect/preview, "
            "or create a new document."
        )


def ensure_docx_suffix(name: str) -> str:
    path = Path(name)
    return name if path.suffix.lower() == ".docx" else str(path.with_suffix(".docx"))

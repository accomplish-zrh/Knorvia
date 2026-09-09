"""XLSX adapter: strict DSL → openpyxl (generated) or narrow patch (imported).

Generated workbooks (created by this runtime) are written with openpyxl —
that is safe because openpyxl produced them. Imported workbooks go through
``xlsx_patch.apply_narrow_patch``; openpyxl never saves an imported file.
"""

from __future__ import annotations

from io import BytesIO
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import range_boundaries
from openpyxl.worksheet.worksheet import Worksheet

from knorvia.services.office_artifacts.adapters.base import AdapterOutcome
from knorvia.services.office_artifacts.adapters.xlsx_patch import apply_narrow_patch
from knorvia.services.office_artifacts.contracts import (
    ClearCellsOp,
    CopyStyleOp,
    FreezePanesOp,
    InvalidOperationError,
    MergeCellsOp,
    SetCellOp,
    SetColumnWidthOp,
    SetFormulaOp,
    SetRangeOp,
    SetRowHeightOp,
    UnmergeCellsOp,
    UnsupportedOperationError,
    cell_name,
    parse_cell,
    parse_range,
    validate_sheet_name,
)


def _openpyxl_value(op: SetCellOp) -> Any:
    if op.text is not None:
        return op.text
    if op.number is not None:
        return op.number
    if op.boolean is not None:
        return op.boolean
    if op.value_date is not None:
        return op.value_date
    if op.value_datetime is not None:
        return op.value_datetime
    return None


def _pick_sheet(wb: Any, sheet: str, *, create_missing: bool) -> Worksheet:
    if not str(sheet or "").strip():
        return wb.active
    if sheet in wb.sheetnames:
        return wb[sheet]
    if create_missing:
        return wb.create_sheet(title=validate_sheet_name(sheet))
    raise InvalidOperationError(
        f"sheet {sheet!r} not found; existing: {', '.join(wb.sheetnames)}"
    )


def create_generated_xlsx(sheet_name: str = "Sheet") -> bytes:
    workbook = Workbook()
    workbook.active.title = validate_sheet_name(sheet_name or "Sheet")
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def apply_to_generated(data: bytes, operations: list[Any]) -> AdapterOutcome:
    """Apply the strict DSL to a runtime-generated workbook via openpyxl."""
    try:
        wb = load_workbook(BytesIO(data))
    except Exception as exc:  # noqa: BLE001
        raise InvalidOperationError(f"generated workbook could not be reopened: {exc}")
    touched = 0
    targets: list[str] = []
    for op in operations:
        ws = _pick_sheet(wb, op.sheet, create_missing=isinstance(op, (SetCellOp, SetFormulaOp)))
        if isinstance(op, SetCellOp):
            ws[op.cell.upper()] = _openpyxl_value(op)
            touched += 1
            targets.append(op.cell.upper())
        elif isinstance(op, SetFormulaOp):
            ws[op.cell.upper()] = op.formula
            touched += 1
            targets.append(op.cell.upper())
        elif isinstance(op, ClearCellsOp):
            min_col, min_row, max_col, max_row = parse_range(op.range)
            for row in range(min_row, max_row + 1):
                for col in range(min_col, max_col + 1):
                    ws.cell(row=row, column=col).value = None
                    touched += 1
                    targets.append(f"{ws.cell(row=row, column=col).coordinate}")
        elif isinstance(op, SetRangeOp):
            min_col, min_row, max_col, max_row = parse_range(op.range)
            for r_offset, row_values in enumerate(op.rows):
                for c_offset, value in enumerate(row_values):
                    cell = ws.cell(row=min_row + r_offset, column=min_col + c_offset)
                    cell.value = value
                    touched += 1
                    targets.append(cell.coordinate)
        elif isinstance(op, CopyStyleOp):
            source_ws = _pick_sheet(wb, op.sheet, create_missing=False)
            s_col, s_row = parse_cell(op.source_cell)
            source = source_ws.cell(row=s_row, column=s_col)
            min_col, min_row, max_col, max_row = parse_range(op.target_range)
            from copy import copy as _copy

            for row in range(min_row, max_row + 1):
                for col in range(min_col, max_col + 1):
                    target = source_ws.cell(row=row, column=col)
                    target.font = _copy(source.font)
                    target.fill = _copy(source.fill)
                    target.border = _copy(source.border)
                    target.number_format = source.number_format
                    target.alignment = _copy(source.alignment)
                    touched += 1
                    targets.append(target.coordinate)
        elif isinstance(op, MergeCellsOp):
            ws.merge_cells(op.range)
            targets.append(op.range)
        elif isinstance(op, UnmergeCellsOp):
            normalized = op.range.upper()
            existing = [str(r) for r in ws.merged_cells.ranges]
            match = next((r for r in existing if r.upper() == normalized), None)
            if match is None:
                raise InvalidOperationError(f"no merged range {op.range} to unmerge")
            ws.unmerge_cells(match)
            targets.append(op.range)
        elif isinstance(op, SetRowHeightOp):
            ws.row_dimensions[op.row].height = op.height
            targets.append(f"row:{op.row}")
        elif isinstance(op, SetColumnWidthOp):
            ws.column_dimensions[op.column].width = op.width
            targets.append(f"col:{op.column}")
        elif isinstance(op, FreezePanesOp):
            ws.freeze_panes = op.cell.upper() if op.cell else None
            targets.append(f"freeze:{op.cell or 'none'}")
        else:  # pragma: no cover
            raise UnsupportedOperationError(
                f"operation {type(op).__name__} is not supported",
                reason="unsupported_operation",
            )
    buffer = BytesIO()
    wb.save(buffer)
    return AdapterOutcome(
        new_bytes=buffer.getvalue(),
        calculation_required=any(isinstance(op, SetFormulaOp) for op in operations),
        touched_cells=touched,
        target_addresses=sorted(set(targets)),
    )


def apply_to_imported(data: bytes, operations: list[Any]) -> AdapterOutcome:
    return apply_narrow_patch(data, operations)


# ----------------------------------------------------------------------
# Legacy actions for generated workbooks (compat facade only — not part
# of the strict DSL, never offered for imported artifacts).
# ----------------------------------------------------------------------


def legacy_action(
    data: bytes, action: str, kwargs: dict[str, Any]
) -> tuple[AdapterOutcome, str]:
    """Run one legacy office_document action against a generated workbook."""
    try:
        wb = load_workbook(BytesIO(data)) if data else None
    except Exception as exc:  # noqa: BLE001
        raise InvalidOperationError(f"workbook could not be reopened: {exc}")
    info = ""
    touched = 0
    targets: list[str] = []
    sheet_name = str(kwargs.get("sheet") or "").strip()

    if action == "add_sheet":
        name = validate_sheet_name(sheet_name)
        if wb is None or name in wb.sheetnames:
            raise InvalidOperationError(f"sheet {name!r} already exists")
        wb.create_sheet(title=name)
        info = f"Added sheet {name!r}."
        targets.append(name)
    elif action in ("write_cells", "formula"):
        if wb is None:
            raise InvalidOperationError("workbook does not exist; create it first")
        ws = _pick_sheet(wb, sheet_name, create_missing=True)
        cells = _coerce_cells(kwargs.get("cells") or kwargs.get("formula_cells"))
        formulas = 0
        written = 0
        for key, value in cells.items():
            if isinstance(value, str) and value.lstrip().startswith("="):
                formulas += 1
            ws[key.upper()] = value
            written += 1
            targets.append(key.upper())
        touched = written
        info = f"Wrote {written} cell(s) ({formulas} formula(s)) on {ws.title!r}."
    elif action == "style":
        if wb is None:
            raise InvalidOperationError("workbook does not exist; create it first")
        ws = _pick_sheet(wb, sheet_name, create_missing=True)
        specs = kwargs.get("styles")
        if not isinstance(specs, list) or not specs:
            raise InvalidOperationError("`styles` must be a non-empty list of style objects")
        applied = _legacy_style(ws, specs)
        touched = applied
        info = f"Applied {len(specs)} style rule(s) to {applied} cell(s)."
    elif action == "chart":
        if wb is None:
            raise InvalidOperationError("workbook does not exist; create it first")
        ws = _pick_sheet(wb, sheet_name, create_missing=True)
        spec = kwargs.get("chart")
        if isinstance(spec, str):
            try:
                import json as _json

                spec = _json.loads(spec)
            except _json.JSONDecodeError:
                spec = None
        if not isinstance(spec, dict):
            raise InvalidOperationError(
                '`chart` is required: {type:"bar|line|pie", data_range:"A1:B5", title:"..."}'
            )
        chart_type = str(spec.get("type") or "bar").strip().lower()
        data_range = str(spec.get("data_range") or "").strip()
        title = str(spec.get("title") or "").strip()
        if chart_type not in ("bar", "line", "pie"):
            raise InvalidOperationError("chart.type must be one of: bar, line, pie")
        if not data_range:
            raise InvalidOperationError("chart.data_range is required (e.g. A1:B5)")
        mapping = {"bar": BarChart, "line": LineChart, "pie": PieChart}
        min_col, min_row, max_col, max_row = range_boundaries(data_range)
        chart = mapping[chart_type]()
        if title:
            chart.title = title
        chart.add_data(Reference(ws, min_col=min_col, min_row=min_row, max_col=max_col, max_row=max_row), titles_from_data=True)
        if chart_type == "pie" and max_col > min_col:
            chart.set_categories(Reference(ws, min_col=min_col, min_row=min_row + 1, max_row=max_row))
        from openpyxl.utils import get_column_letter

        ws.add_chart(chart, f"{get_column_letter(max_col + 2)}{min_row}")
        info = f"Added {chart_type} chart ({data_range})."
        targets.append(data_range)
    else:
        raise InvalidOperationError(f"legacy action {action!r} is not supported here")

    buffer = BytesIO()
    wb.save(buffer)
    outcome = AdapterOutcome(
        new_bytes=buffer.getvalue(),
        calculation_required=any(
            isinstance(v, str) and str(v).lstrip().startswith("=")
            for op_cells in [kwargs.get("cells") or kwargs.get("formula_cells") or {}]
            for v in (op_cells.values() if isinstance(op_cells, dict) else [])
        ),
        touched_cells=touched,
        target_addresses=sorted(set(targets)),
    )
    return outcome, info


def _coerce_cells(raw: Any) -> dict[str, Any]:
    import json as _json

    if isinstance(raw, str):
        try:
            raw = _json.loads(raw)
        except _json.JSONDecodeError:
            pass
    if isinstance(raw, list):
        # Legacy 2D grid [["a","b"],[1,2]] anchored at A1.
        grid: dict[str, Any] = {}
        for r_offset, row in enumerate(raw):
            if not isinstance(row, (list, tuple)):
                raise InvalidOperationError("cells 2D array must be a list of rows")
            for c_offset, value in enumerate(row):
                grid[cell_name(1 + c_offset, 1 + r_offset)] = value
        if not grid:
            raise InvalidOperationError("cells must not be empty")
        return grid
    if not isinstance(raw, dict) or not raw:
        raise InvalidOperationError('cells must be an object {"A1": value, ...} or a 2D array')
    result: dict[str, Any] = {}
    for key, value in raw.items():
        if isinstance(value, str) and len(value) > 32767:
            raise InvalidOperationError("cell text exceeds 32767 characters")
        result[str(key)] = value
    return result


def _legacy_style(ws: Worksheet, specs: list[Any]) -> int:
    applied = 0
    for spec in specs:
        if not isinstance(spec, dict):
            raise InvalidOperationError("each style entry must be an object")
        target = str(spec.get("target") or "").strip()
        if not target:
            raise InvalidOperationError("style.target is required (e.g. A1:D1)")
        min_col, min_row, max_col, max_row = range_boundaries(target)
        font_kwargs: dict[str, Any] = {}
        if "bold" in spec:
            font_kwargs["bold"] = bool(spec.get("bold"))
        if spec.get("color"):
            font_kwargs["color"] = _rgb(str(spec["color"]))
        if spec.get("font_size") is not None:
            font_kwargs["size"] = float(spec["font_size"])
        fill = PatternFill(fill_type="solid", fgColor=_rgb(str(spec["bg"]))) if spec.get("bg") else None
        font = Font(**font_kwargs) if font_kwargs else None
        for row in ws.iter_rows(min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col):
            for cell in row:
                if font is not None:
                    cell.font = font
                if fill is not None:
                    cell.fill = fill
                applied += 1
    return applied


def _rgb(value: str) -> str:
    raw = str(value or "").strip().lstrip("#")
    if len(raw) == 6 and all(ch in "0123456789abcdefABCDEF" for ch in raw):
        return "FF" + raw.upper()
    if len(raw) == 8 and all(ch in "0123456789abcdefABCDEF" for ch in raw):
        return raw.upper()
    raise InvalidOperationError(f"invalid colour {value!r}; use #RRGGBB")


def readback_cells(data: bytes, sheet: str, targets: list[str]) -> dict[str, dict[str, Any]]:
    from knorvia.services.office_artifacts.adapters.xlsx_reader import XlsxReader

    reader = XlsxReader(data)
    return reader.snapshot_cells(sheet, targets)


def snapshot_for_diff(data: bytes, sheet: str, targets: list[str]) -> dict[str, dict[str, Any]]:
    return readback_cells(data, sheet, targets)


__all__ = [
    "apply_to_generated",
    "apply_to_imported",
    "create_generated_xlsx",
    "legacy_action",
    "readback_cells",
    "snapshot_for_diff",
]

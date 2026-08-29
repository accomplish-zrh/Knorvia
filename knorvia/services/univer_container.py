"""ZIP-backed ``.univer`` multi-unit Office container.

A ``.univer`` file is an OOXML-style ZIP containing:

* ``manifest.json`` — ``units`` array
  (``{id, type: sheet|doc|slide, file, name?}``) plus optional ``refs``
* ``units/<id>.xlsx|docx|pptx`` — each unit's native Office payload

Cross-unit chart linkage is out of scope for this tier. ``refs`` records the
relationship (``from`` → ``to`` + ``range``), and Slide units may embed a
text placeholder such as ``数据来源: Sheet!A1:B5`` so the intent is visible
when the deck is opened outside Knorvia.
"""

from __future__ import annotations

from collections.abc import Mapping, Sequence
from io import BytesIO
import json
from pathlib import Path
import re
from typing import Any, Literal
import uuid
import zipfile

UnitType = Literal["sheet", "doc", "slide"]

UNIT_TYPES: tuple[UnitType, ...] = ("sheet", "doc", "slide")
UNIT_SUFFIX: dict[str, str] = {
    "sheet": ".xlsx",
    "doc": ".docx",
    "slide": ".pptx",
}
MANIFEST_NAME = "manifest.json"
UNITS_DIR = "units"
CONTAINER_VERSION = 1
_ID_RE = re.compile(r"^[A-Za-z][A-Za-z0-9_-]{0,31}$")


class ContainerError(ValueError):
    """Raised when a ``.univer`` container is missing, corrupt, or invalid."""


def _as_path(path: str | Path) -> Path:
    return Path(path).expanduser().resolve()


def _validate_unit_type(unit_type: str) -> UnitType:
    value = str(unit_type or "").strip().lower()
    if value not in UNIT_TYPES:
        raise ContainerError(f"unit type must be one of {', '.join(UNIT_TYPES)}; got {unit_type!r}")
    return value  # type: ignore[return-value]


def _validate_unit_id(unit_id: str) -> str:
    value = str(unit_id or "").strip()
    if not _ID_RE.fullmatch(value):
        raise ContainerError(
            f"invalid unit id {unit_id!r}; use letters/digits/_/- starting with a letter"
        )
    return value


def _new_unit_id(unit_type: UnitType, existing: set[str]) -> str:
    base = {"sheet": "sheet", "doc": "doc", "slide": "slide"}[unit_type]
    if base not in existing:
        return base
    for index in range(2, 1000):
        candidate = f"{base}{index}"
        if candidate not in existing:
            return candidate
    return f"{base}_{uuid.uuid4().hex[:6]}"


def _unit_archive_path(unit_id: str, unit_type: UnitType) -> str:
    return f"{UNITS_DIR}/{unit_id}{UNIT_SUFFIX[unit_type]}"


def _empty_manifest(
    units: list[dict[str, Any]] | None = None,
    refs: list[dict[str, Any]] | None = None,
) -> dict[str, Any]:
    return {
        "version": CONTAINER_VERSION,
        "units": list(units or []),
        "refs": list(refs or []),
    }


def _normalise_ref(raw: Mapping[str, Any], *, default_from: str | None = None) -> dict[str, Any]:
    source = str(raw.get("from") or default_from or "").strip()
    target = str(raw.get("to") or raw.get("target") or "").strip()
    data_range = str(raw.get("range") or raw.get("data_range") or "").strip()
    kind = str(raw.get("kind") or "data_source").strip() or "data_source"
    if not source or not target:
        raise ContainerError("each ref needs `from` and `to` unit ids")
    if not data_range:
        raise ContainerError("each ref needs a `range` (e.g. A1:B5)")
    return {
        "from": source,
        "to": target,
        "range": data_range,
        "kind": kind,
    }


def _read_manifest(zf: zipfile.ZipFile) -> dict[str, Any]:
    try:
        raw = zf.read(MANIFEST_NAME)
    except KeyError as exc:
        raise ContainerError(f"missing {MANIFEST_NAME} in .univer container") from exc
    try:
        data = json.loads(raw.decode("utf-8"))
    except (UnicodeDecodeError, json.JSONDecodeError) as exc:
        raise ContainerError(f"corrupt {MANIFEST_NAME}: {exc}") from exc
    if not isinstance(data, dict):
        raise ContainerError("manifest.json must be a JSON object")
    units = data.get("units")
    if not isinstance(units, list):
        raise ContainerError("manifest.units must be an array")
    refs = data.get("refs") if isinstance(data.get("refs"), list) else []
    return _empty_manifest(
        [row for row in units if isinstance(row, dict)],
        [row for row in refs if isinstance(row, dict)],
    )


def _write_zip(path: Path, manifest: dict[str, Any], files: Mapping[str, bytes]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_suffix(path.suffix + ".tmp")
    try:
        with zipfile.ZipFile(tmp, "w", compression=zipfile.ZIP_DEFLATED) as zf:
            zf.writestr(
                MANIFEST_NAME,
                json.dumps(manifest, ensure_ascii=False, indent=2) + "\n",
            )
            for name, payload in files.items():
                zf.writestr(name, payload)
        tmp.replace(path)
    finally:
        if tmp.exists():
            tmp.unlink(missing_ok=True)


def _load_zip_files(path: Path) -> tuple[dict[str, Any], dict[str, bytes]]:
    if not path.is_file():
        raise ContainerError(f"container not found: {path}")
    try:
        with zipfile.ZipFile(path, "r") as zf:
            if zf.testzip() is not None:
                raise ContainerError(f"corrupt zip archive: {path.name}")
            manifest = _read_manifest(zf)
            files: dict[str, bytes] = {}
            for info in zf.infolist():
                if info.is_dir() or info.filename == MANIFEST_NAME:
                    continue
                files[info.filename] = zf.read(info.filename)
            return manifest, files
    except zipfile.BadZipFile as exc:
        raise ContainerError(f"corrupt zip archive: {path.name}") from exc


def _make_empty_unit_bytes(unit_type: UnitType, *, title: str = "") -> bytes:
    """Build a minimal native Office file for an empty unit."""
    if unit_type == "sheet":
        from openpyxl import Workbook

        workbook = Workbook()
        ws = workbook.active
        ws.title = (title or "Sheet")[:31] or "Sheet"
        buf = BytesIO()
        workbook.save(buf)
        return buf.getvalue()

    if unit_type == "doc":
        from docx import Document

        document = Document()
        if title:
            document.add_heading(title, level=1)
        else:
            document.add_paragraph("")
        buf = BytesIO()
        document.save(buf)
        return buf.getvalue()

    from pptx import Presentation
    from pptx.util import Inches

    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)
    layout = presentation.slide_layouts[0]
    slide = presentation.slides.add_slide(layout)
    if slide.shapes.title is not None:
        slide.shapes.title.text = title or "Slide"
    buf = BytesIO()
    presentation.save(buf)
    return buf.getvalue()


def _append_slide_placeholder(pptx_bytes: bytes, placeholder: str) -> bytes:
    """Stamp a visible data-source note onto the first slide."""
    from pptx import Presentation
    from pptx.util import Inches, Pt

    presentation = Presentation(BytesIO(pptx_bytes))
    if not presentation.slides:
        layout = presentation.slide_layouts[5]
        presentation.slides.add_slide(layout)
    slide = presentation.slides[0]
    box = slide.shapes.add_textbox(Inches(0.5), Inches(6.6), Inches(12.3), Inches(0.6))
    frame = box.text_frame
    frame.clear()
    paragraph = frame.paragraphs[0]
    run = paragraph.add_run()
    run.text = placeholder
    run.font.size = Pt(12)
    buf = BytesIO()
    presentation.save(buf)
    return buf.getvalue()


def _placeholder_for_refs(unit_id: str, refs: Sequence[Mapping[str, Any]]) -> str | None:
    notes: list[str] = []
    for raw in refs:
        if str(raw.get("from") or unit_id) != unit_id:
            continue
        target = str(raw.get("to") or "").strip()
        data_range = str(raw.get("range") or raw.get("data_range") or "").strip()
        if target and data_range:
            notes.append(f"数据来源: {target}!{data_range}")
    if not notes:
        return None
    return " · ".join(notes)


def pack_univer(
    manifest_entries: Sequence[Mapping[str, Any]],
    out_path: str | Path,
    *,
    refs: Sequence[Mapping[str, Any]] | None = None,
) -> dict[str, Any]:
    """Create a ``.univer`` ZIP from unit descriptors.

    Each entry needs ``id``, ``type``, and either ``source`` (``Path`` /
    bytes / readable path string) or an empty native file is generated.
    Optional ``name`` is stored on the manifest for UI labels.

    Returns the written manifest.
    """
    path = _as_path(out_path)
    if path.suffix.lower() != ".univer":
        path = path.with_suffix(".univer")

    units: list[dict[str, Any]] = []
    files: dict[str, bytes] = {}
    seen: set[str] = set()

    for raw in manifest_entries:
        unit_type = _validate_unit_type(str(raw.get("type") or ""))
        unit_id = _validate_unit_id(str(raw.get("id") or _new_unit_id(unit_type, seen)))
        if unit_id in seen:
            raise ContainerError(f"duplicate unit id {unit_id!r}")
        seen.add(unit_id)
        name = str(raw.get("name") or unit_id).strip() or unit_id
        archive_name = str(raw.get("file") or _unit_archive_path(unit_id, unit_type))
        source = raw.get("source")
        if source is None:
            payload = _make_empty_unit_bytes(unit_type, title=name)
        elif isinstance(source, (bytes, bytearray)):
            payload = bytes(source)
        else:
            source_path = _as_path(source)
            if not source_path.is_file():
                raise ContainerError(f"unit source not found: {source_path}")
            payload = source_path.read_bytes()
        units.append(
            {
                "id": unit_id,
                "type": unit_type,
                "file": archive_name,
                "name": name,
            }
        )
        files[archive_name] = payload

    normalised_refs = [_normalise_ref(row) for row in (refs or ())]
    # Apply slide placeholders for declared refs.
    for ref in normalised_refs:
        for unit in units:
            if unit["id"] != ref["from"] or unit["type"] != "slide":
                continue
            note = _placeholder_for_refs(unit["id"], [ref])
            if note:
                files[unit["file"]] = _append_slide_placeholder(files[unit["file"]], note)

    manifest = _empty_manifest(units, normalised_refs)
    _write_zip(path, manifest, files)
    return manifest


def open_univer(path: str | Path) -> dict[str, Any]:
    """Parse and validate the container manifest (does not extract units)."""
    manifest, files = _load_zip_files(_as_path(path))
    for unit in manifest["units"]:
        archive_name = str(unit.get("file") or "")
        if not archive_name or archive_name not in files:
            raise ContainerError(f"manifest references missing unit file {archive_name!r}")
    return manifest


def list_units(path: str | Path) -> list[dict[str, Any]]:
    """Return the ``units`` array from an existing container."""
    return list(open_univer(path).get("units") or [])


def add_unit(
    path: str | Path,
    unit_type: str,
    name: str | None = None,
    source_bytes: bytes | None = None,
    *,
    unit_id: str | None = None,
    refs: Sequence[Mapping[str, Any]] | None = None,
) -> dict[str, Any]:
    """Append one unit to an existing ``.univer`` container.

    ``refs`` entries may omit ``from`` (defaults to the new unit id). When
    the new unit is a Slide and refs point at a Sheet range, a text
    placeholder is stamped onto the first slide.
    """
    container = _as_path(path)
    manifest, files = _load_zip_files(container)
    validated_type = _validate_unit_type(unit_type)
    existing_ids = {str(unit.get("id") or "") for unit in manifest["units"] if unit.get("id")}
    new_id = _validate_unit_id(unit_id) if unit_id else _new_unit_id(validated_type, existing_ids)
    if new_id in existing_ids:
        raise ContainerError(f"unit id {new_id!r} already exists")

    label = str(name or new_id).strip() or new_id
    archive_name = _unit_archive_path(new_id, validated_type)
    payload = (
        bytes(source_bytes)
        if source_bytes is not None
        else _make_empty_unit_bytes(validated_type, title=label)
    )

    normalised_refs = [_normalise_ref(row, default_from=new_id) for row in (refs or ())]
    if validated_type == "slide":
        note = _placeholder_for_refs(new_id, normalised_refs)
        if note:
            payload = _append_slide_placeholder(payload, note)

    unit = {
        "id": new_id,
        "type": validated_type,
        "file": archive_name,
        "name": label,
    }
    manifest["units"].append(unit)
    manifest["refs"].extend(normalised_refs)
    files[archive_name] = payload
    _write_zip(container, manifest, files)
    return unit


def remove_unit(path: str | Path, unit_id: str) -> dict[str, Any]:
    """Remove a unit (and refs that mention it) from the container."""
    container = _as_path(path)
    target = _validate_unit_id(unit_id)
    manifest, files = _load_zip_files(container)
    remaining: list[dict[str, Any]] = []
    removed: dict[str, Any] | None = None
    for unit in manifest["units"]:
        if str(unit.get("id") or "") == target:
            removed = unit
            archive_name = str(unit.get("file") or "")
            files.pop(archive_name, None)
            continue
        remaining.append(unit)
    if removed is None:
        raise ContainerError(f"unit {target!r} not found")
    manifest["units"] = remaining
    manifest["refs"] = [
        ref
        for ref in manifest["refs"]
        if str(ref.get("from") or "") != target and str(ref.get("to") or "") != target
    ]
    _write_zip(container, manifest, files)
    return removed


def read_unit_bytes(path: str | Path, unit_id: str) -> tuple[bytes, dict[str, Any]]:
    """Return ``(bytes, unit_manifest_entry)`` for one unit id."""
    container = _as_path(path)
    target = _validate_unit_id(unit_id)
    manifest, files = _load_zip_files(container)
    for unit in manifest["units"]:
        if str(unit.get("id") or "") != target:
            continue
        archive_name = str(unit.get("file") or "")
        payload = files.get(archive_name)
        if payload is None:
            raise ContainerError(f"unit file missing for {target!r}")
        return payload, unit
    raise ContainerError(f"unit {target!r} not found")


def export_unit(path: str | Path, unit_id: str, out_path: str | Path) -> Path:
    """Write one unit's native Office file to ``out_path`` and return it."""
    payload, unit = read_unit_bytes(path, unit_id)
    dest = _as_path(out_path)
    unit_type = _validate_unit_type(str(unit.get("type") or "sheet"))
    if dest.suffix.lower() != UNIT_SUFFIX[unit_type]:
        dest = dest.with_suffix(UNIT_SUFFIX[unit_type])
    dest.parent.mkdir(parents=True, exist_ok=True)
    dest.write_bytes(payload)
    return dest


def unit_media_type(unit_type: str) -> str:
    """MIME type for a unit's native Office payload."""
    mapping = {
        "sheet": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "doc": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "slide": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    }
    return mapping.get(_validate_unit_type(unit_type), "application/octet-stream")

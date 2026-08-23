"""Encode and extract personal-library Word/Excel files.

Creates real ``.docx`` / ``.xlsx`` bytes and reads them back through the
shipped ``document_extractor`` so Agent, UI, and tests share one path.
Legacy ``.doc`` / ``.xls`` are not encoded here.
"""

from __future__ import annotations

import io
import re

WORD_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
EXCEL_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
WORD_EXT = ".docx"
EXCEL_EXT = ".xlsx"

_SHEET_BANNER = re.compile(r"^--- Sheet: [^\n]* ---\n?", re.MULTILINE)


def office_ext(kind: str) -> str:
    if kind == "word":
        return WORD_EXT
    if kind == "excel":
        return EXCEL_EXT
    raise ValueError("Unsupported office kind")


def office_mime(kind: str) -> str:
    if kind == "word":
        return WORD_MIME
    if kind == "excel":
        return EXCEL_MIME
    raise ValueError("Unsupported office kind")


def encode_library_office(kind: str, text: str) -> bytes:
    body = str(text or "")
    if kind == "word":
        return _encode_docx(body)
    if kind == "excel":
        return _encode_xlsx(body)
    raise ValueError("Unsupported office kind")


def extract_library_office(kind: str, data: bytes) -> str:
    """Extract the written body from a library Word/Excel file.

    Uses ``extract_text_from_bytes`` (not a parallel parser). Excel sheet
    banners added by that extractor are stripped so a single-cell body
    round-trips to the same string.
    """
    from knorvia.utils.document_extractor import EmptyDocumentError, extract_text_from_bytes

    if not data:
        return ""
    if kind not in {"word", "excel"}:
        raise ValueError("Unsupported office kind")
    filename = f"library{office_ext(kind)}"
    try:
        raw = extract_text_from_bytes(filename, data)
    except EmptyDocumentError:
        return ""
    if kind == "excel":
        raw = _SHEET_BANNER.sub("", raw)
    return raw


def _encode_docx(text: str) -> bytes:
    from docx import Document

    document = Document()
    if document.paragraphs:
        document.paragraphs[0].text = text
    else:
        document.add_paragraph(text)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _encode_xlsx(text: str) -> bytes:
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sheet"
    if text:
        sheet["A1"] = text
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()

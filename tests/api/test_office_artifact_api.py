"""Typed OAR v2 endpoints under ``/api/v1/chat/office-drafts``."""

from __future__ import annotations

from collections.abc import Callable
from pathlib import Path

from fastapi import FastAPI
from fastapi.testclient import TestClient
import pytest

from knorvia.services.auth import TokenPayload
from knorvia.services.office_artifacts.service import OfficeArtifactService
from knorvia.services.path_service import PathService

AppFactory = Callable[[dict[str, TokenPayload | None], bool], tuple[TestClient, Path, Path]]


@pytest.fixture
def artifact_app(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> AppFactory:
    from knorvia.api.routers import auth as auth_router
    from knorvia.api.routers import office_drafts
    from knorvia.multi_user import paths as multi_user_paths

    admin_root = tmp_path / "data"
    users_root = admin_root / "users"
    monkeypatch.setattr(multi_user_paths, "ADMIN_WORKSPACE_ROOT", admin_root)
    monkeypatch.setattr(multi_user_paths, "USERS_ROOT", users_root)
    monkeypatch.setattr(multi_user_paths, "_path_services", {})

    def make_app(
        tokens: dict[str, TokenPayload | None], auth_enabled: bool = True
    ) -> tuple[TestClient, Path, Path]:
        monkeypatch.setattr(auth_router, "AUTH_ENABLED", auth_enabled)
        monkeypatch.setattr(auth_router, "decode_token", lambda token: tokens.get(token))
        app = FastAPI()
        app.include_router(office_drafts.router, prefix="/api/v1")
        return TestClient(app), admin_root, users_root

    return make_app


def _alice_v2_draft(users_root: Path) -> tuple[PathService, OfficeArtifactService, str, str]:
    service = PathService(workspace_root=users_root / "u_alice")
    task_dir = service.get_task_workspace("chat", "turn-1")
    runtime = OfficeArtifactService(
        task_dir=task_dir,
        workspace_dir=task_dir / "exec",
        public_root=service.get_public_outputs_root(),
    )
    payload = runtime.create_generated("report.xlsx", "xlsx", sheet="Data")
    return service, runtime, payload["draft_id"], payload["artifact"]["artifact_id"]


def test_get_card_serves_v2_draft(artifact_app) -> None:
    alice = TokenPayload(username="alice", role="user", user_id="u_alice")
    client, _admin, users_root = artifact_app({"alice-token": alice})
    _service, _runtime, draft_id, artifact_id = _alice_v2_draft(users_root)

    with client:
        client.cookies.set("dt_token", "alice-token")
        response = client.get(f"/api/v1/chat/office-drafts/{draft_id}")

    assert response.status_code == 200
    body = response.json()
    assert body["draft_id"] == draft_id
    assert body["status"] == "draft"
    assert body["artifacts"][0]["artifact_id"] == artifact_id


def test_overview_and_range_reads(artifact_app) -> None:
    alice = TokenPayload(username="alice", role="user", user_id="u_alice")
    client, _admin, users_root = artifact_app({"alice-token": alice})
    _service, _runtime, draft_id, artifact_id = _alice_v2_draft(users_root)

    with client:
        client.cookies.set("dt_token", "alice-token")
        overview = client.get(
            f"/api/v1/chat/office-drafts/{draft_id}/artifacts/{artifact_id}/overview"
        )
        rng = client.get(
            f"/api/v1/chat/office-drafts/{draft_id}/artifacts/{artifact_id}/range",
            params={"sheet": "Data", "range": "A1:B2"},
        )

    assert overview.status_code == 200
    assert overview.json()["current_revision"] == 0
    assert rng.status_code == 200
    assert rng.json()["revision"] == 0


def test_user_operation_batch_applies_and_conflicts(artifact_app) -> None:
    alice = TokenPayload(username="alice", role="user", user_id="u_alice")
    client, _admin, users_root = artifact_app({"alice-token": alice})
    _service, _runtime, draft_id, artifact_id = _alice_v2_draft(users_root)
    url = f"/api/v1/chat/office-drafts/{draft_id}/artifacts/{artifact_id}/operations"
    batch = {
        "base_revision": 0,
        "operations": [{"op": "set_cell", "sheet": "Data", "cell": "A1", "number": 42}],
    }

    with client:
        client.cookies.set("dt_token", "alice-token")
        ok = client.post(url, json=batch)
        conflict = client.post(url, json=batch)

    assert ok.status_code == 200
    payload = ok.json()
    assert payload["revision_after"] == 1
    verification = payload["verification"]
    assert all(
        verification[key]
        for key in ("reopened", "target_readback", "zip_structure_valid", "untouched_parts_verified")
    )
    assert conflict.status_code == 409


def test_invalid_operation_batch_is_422(artifact_app) -> None:
    alice = TokenPayload(username="alice", role="user", user_id="u_alice")
    client, _admin, users_root = artifact_app({"alice-token": alice})
    _service, _runtime, draft_id, artifact_id = _alice_v2_draft(users_root)
    url = f"/api/v1/chat/office-drafts/{draft_id}/artifacts/{artifact_id}/operations"

    with client:
        client.cookies.set("dt_token", "alice-token")
        response = client.post(
            url,
            json={
                "base_revision": 0,
                "operations": [{"op": "set_cell", "sheet": "Data", "cell": "A1"}],
            },
        )

    assert response.status_code == 422


def test_content_download_and_undo(artifact_app) -> None:
    alice = TokenPayload(username="alice", role="user", user_id="u_alice")
    client, _admin, users_root = artifact_app({"alice-token": alice})
    _service, _runtime, draft_id, artifact_id = _alice_v2_draft(users_root)
    base = f"/api/v1/chat/office-drafts/{draft_id}/artifacts/{artifact_id}"

    with client:
        client.cookies.set("dt_token", "alice-token")
        client.post(
            f"{base}/operations",
            json={
                "base_revision": 0,
                "operations": [{"op": "set_cell", "sheet": "Data", "cell": "A1", "number": 7}],
            },
        )
        content = client.get(f"{base}/content")
        undone = client.post(f"{base}/undo")

    assert content.status_code == 200
    assert content.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert undone.status_code == 200
    assert undone.json()["current_revision"] == 0
    assert undone.json()["cursor"] == 0
    assert undone.json()["history_length"] == 2


def test_patch_merge_writes_v2_draft_to_exec(artifact_app) -> None:
    alice = TokenPayload(username="alice", role="user", user_id="u_alice")
    client, _admin, users_root = artifact_app({"alice-token": alice})
    service, _runtime, draft_id, artifact_id = _alice_v2_draft(users_root)
    base = f"/api/v1/chat/office-drafts/{draft_id}/artifacts/{artifact_id}"

    with client:
        client.cookies.set("dt_token", "alice-token")
        client.post(
            f"{base}/operations",
            json={
                "base_revision": 0,
                "operations": [{"op": "set_cell", "sheet": "Data", "cell": "A1", "number": 5}],
            },
        )
        merged = client.patch(
            f"/api/v1/chat/office-drafts/{draft_id}", json={"action": "merge"}
        )

    assert merged.status_code == 200
    assert merged.json()["status"] == "merged"
    official = service.get_task_workspace("chat", "turn-1") / "exec" / "report.xlsx"
    assert official.is_file()


def _library_store(tmp_path: Path):
    from knorvia.services.creative_library.store import CreativeLibraryStore

    store = CreativeLibraryStore(tmp_path / "library")
    return store


def _patch_library_store(monkeypatch: pytest.MonkeyPatch, store) -> None:
    monkeypatch.setattr(
        "knorvia.services.creative_library.store.get_creative_library_store",
        lambda: store,
    )


def test_library_open_apply_merge_writes_entry_back(
    artifact_app, monkeypatch, tmp_path
) -> None:
    import io

    store = _library_store(tmp_path)
    _patch_library_store(monkeypatch, store)
    entry = store.create_entry(kind="excel", title="Budget", content="original")
    entry_id = entry["id"]
    alice = TokenPayload(username="alice", role="user", user_id="u_alice")
    client, _admin, _users_root = artifact_app({"alice-token": alice})

    with client:
        client.cookies.set("dt_token", "alice-token")
        opened = client.post(
            "/api/v1/chat/office-drafts/open",
            json={"source": f"library:{entry_id}"},
        )
        assert opened.status_code == 200, opened.text
        card = opened.json()["card"]
        draft_id = card["draft_id"]
        artifact = card["artifacts"][0]
        base = (
            f"/api/v1/chat/office-drafts/{draft_id}"
            f"/artifacts/{artifact['artifact_id']}"
        )
        overview = client.get(f"{base}/overview")
        assert overview.status_code == 200
        sheet_name = overview.json()["sheets"][0]["name"]
        applied = client.post(
            f"{base}/operations",
            json={
                "base_revision": artifact["current_revision"],
                "operations": [
                    {
                        "op": "set_cell",
                        "sheet": sheet_name,
                        "cell": "A1",
                        "text": "human edit",
                    }
                ],
            },
        )
        assert applied.status_code == 200, applied.text
        merged = client.patch(
            f"/api/v1/chat/office-drafts/{draft_id}", json={"action": "merge"}
        )
    assert merged.status_code == 200, merged.text
    assert merged.json()["status"] == "merged"

    from openpyxl import load_workbook

    payload = store.entry_bytes(entry_id)
    assert payload is not None
    workbook = load_workbook(io.BytesIO(payload[0]))
    worksheet = workbook[workbook.sheetnames[0]]
    assert worksheet["A1"].value == "human edit"


def test_library_merge_conflicts_when_entry_changed(
    artifact_app, monkeypatch, tmp_path
) -> None:
    store = _library_store(tmp_path)
    _patch_library_store(monkeypatch, store)
    entry = store.create_entry(kind="excel", title="Budget", content="original")
    entry_id = entry["id"]
    other = store.create_entry(kind="excel", title="Other", content="other")
    replacement = store.entry_bytes(other["id"])
    assert replacement is not None
    alice = TokenPayload(username="alice", role="user", user_id="u_alice")
    client, _admin, _users_root = artifact_app({"alice-token": alice})

    with client:
        client.cookies.set("dt_token", "alice-token")
        opened = client.post(
            "/api/v1/chat/office-drafts/open",
            json={"source": f"library:{entry_id}"},
        )
        assert opened.status_code == 200, opened.text
        draft_id = opened.json()["card"]["draft_id"]
        # Someone else saves to the entry while the human editor is open.
        stored = store.replace_entry_bytes(entry_id, replacement[0])
        assert stored is not None
        merged = client.patch(
            f"/api/v1/chat/office-drafts/{draft_id}", json={"action": "merge"}
        )
    assert merged.status_code == 409
    original = store.entry_bytes(entry_id)
    assert original is not None
    assert original[0] == replacement[0]


def _library_entry_with_bytes(store, title: str = "Budget"):
    entry = store.create_entry(kind="excel", title=title, content="original")
    entry_id = entry["id"]
    payload = store.entry_bytes(entry_id)
    assert payload is not None
    return entry_id, payload[0]


def test_open_with_matching_expected_hash_succeeds(artifact_app, monkeypatch, tmp_path) -> None:
    import hashlib

    store = _library_store(tmp_path)
    _patch_library_store(monkeypatch, store)
    entry_id, data = _library_entry_with_bytes(store)
    alice = TokenPayload(username="alice", role="user", user_id="u_alice")
    client, _admin, _users_root = artifact_app({"alice-token": alice})

    with client:
        client.cookies.set("dt_token", "alice-token")
        opened = client.post(
            "/api/v1/chat/office-drafts/open",
            json={
                "source": f"library:{entry_id}",
                "expected_base_hash": hashlib.sha256(data).hexdigest(),
            },
        )

    assert opened.status_code == 200, opened.text
    artifact = opened.json()["card"]["artifacts"][0]
    assert artifact["current_hash"] == hashlib.sha256(data).hexdigest()


def test_open_with_stale_expected_hash_conflicts_before_creating_a_draft(
    artifact_app, monkeypatch, tmp_path
) -> None:
    store = _library_store(tmp_path)
    _patch_library_store(monkeypatch, store)
    entry_id, _data = _library_entry_with_bytes(store)
    other = store.create_entry(kind="excel", title="Other", content="other")
    replacement = store.entry_bytes(other["id"])
    assert replacement is not None
    alice = TokenPayload(username="alice", role="user", user_id="u_alice")
    client, _admin, _users_root = artifact_app({"alice-token": alice})

    with client:
        client.cookies.set("dt_token", "alice-token")
        # The editor is still showing the version it loaded; a concurrent
        # save replaced the entry underneath it.
        store.replace_entry_bytes(entry_id, replacement[0])
        opened = client.post(
            "/api/v1/chat/office-drafts/open",
            json={
                "source": f"library:{entry_id}",
                "expected_base_hash": "0" * 64,
            },
        )
        drafts = client.get("/api/v1/chat/office-drafts/does-not-exist")

    assert opened.status_code == 409
    assert "reopen" in opened.json()["detail"].lower()
    assert drafts.status_code == 404


def test_merge_reports_the_published_hash_so_the_editor_can_reanchor(
    artifact_app, monkeypatch, tmp_path
) -> None:
    import hashlib

    store = _library_store(tmp_path)
    _patch_library_store(monkeypatch, store)
    entry_id, _data = _library_entry_with_bytes(store)
    alice = TokenPayload(username="alice", role="user", user_id="u_alice")
    client, _admin, _users_root = artifact_app({"alice-token": alice})

    with client:
        client.cookies.set("dt_token", "alice-token")
        opened = client.post(
            "/api/v1/chat/office-drafts/open",
            json={"source": f"library:{entry_id}"},
        )
        card = opened.json()["card"]
        draft_id = card["draft_id"]
        artifact = card["artifacts"][0]
        base = (
            f"/api/v1/chat/office-drafts/{draft_id}"
            f"/artifacts/{artifact['artifact_id']}"
        )
        overview = client.get(f"{base}/overview")
        sheet_name = overview.json()["sheets"][0]["name"]
        applied = client.post(
            f"{base}/operations",
            json={
                "base_revision": artifact["current_revision"],
                "operations": [
                    {
                        "op": "set_cell",
                        "sheet": sheet_name,
                        "cell": "B4",
                        "number": 42,
                    }
                ],
            },
        )
        assert applied.status_code == 200, applied.text
        merged = client.patch(
            f"/api/v1/chat/office-drafts/{draft_id}", json={"action": "merge"}
        )

    assert merged.status_code == 200, merged.text
    stored = store.entry_bytes(entry_id)
    assert stored is not None
    published = merged.json()["artifacts"][0]["current_hash"]
    assert published == hashlib.sha256(stored[0]).hexdigest()


def test_diff_endpoint_counts_entries_not_result_keys(artifact_app) -> None:
    alice = TokenPayload(username="alice", role="user", user_id="u_alice")
    client, _admin, users_root = artifact_app({"alice-token": alice})
    _service, _runtime, draft_id, artifact_id = _alice_v2_draft(users_root)
    base = f"/api/v1/chat/office-drafts/{draft_id}/artifacts/{artifact_id}"

    with client:
        client.cookies.set("dt_token", "alice-token")
        overview = client.get(f"{base}/overview")
        sheet_name = overview.json()["sheets"][0]["name"]
        applied = client.post(
            f"{base}/operations",
            json={
                "base_revision": 0,
                "operations": [
                    {"op": "set_cell", "sheet": sheet_name, "cell": "A1", "text": "one"},
                    {"op": "set_cell", "sheet": sheet_name, "cell": "B2", "number": 2},
                    {"op": "set_cell", "sheet": sheet_name, "cell": "C3", "boolean": True},
                ],
            },
        )
        assert applied.status_code == 200, applied.text
        diff = client.get(
            f"{base}/diff", params={"from_revision": 0, "to_revision": 1}
        )

    assert diff.status_code == 200, diff.text
    body = diff.json()
    assert body["entry_count"] == 3
    assert body["entry_count"] == len(body["diff"]["entries"])
    assert body["omitted_count"] == 0

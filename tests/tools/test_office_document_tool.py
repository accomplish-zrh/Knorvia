"""Unit tests for the structured ``office_document`` tool."""

from __future__ import annotations

from pathlib import Path
from typing import Any

from knorvia.agents._shared.tool_composition import ToolMountFlags, compose_enabled_tools
from knorvia.core.tool_protocol import ToolResult
from knorvia.tools.builtin import BUILTIN_TOOL_NAMES, CONFIGURABLE_BUILTIN_TOOL_NAMES
from knorvia.tools.office_document import ACTIONS, OfficeDocumentTool, execute_office_document


class _EmptyRegistry:
    @staticmethod
    def get_enabled(_selected: list[str]) -> list[Any]:
        return []


def _run(workspace: Path, **kwargs: Any) -> ToolResult:
    return execute_office_document(
        {
            "_workspace_dir": str(workspace),
            "_sandbox_workdir": str(workspace),
            **kwargs,
        }
    )


def _xlsx(path: Path):
    from openpyxl import load_workbook

    return load_workbook(path)


def test_definition_lists_actions_and_teaches_the_flow() -> None:
    definition = OfficeDocumentTool().get_definition()
    assert definition.name == "office_document"
    assert "create →" in definition.description or "create ->" in definition.description
    schema = definition.to_openai_schema()["function"]["parameters"]
    assert schema["required"] == ["action"]
    assert set(schema["properties"]["action"]["enum"]) == set(ACTIONS)


def test_create_write_read_round_trip(tmp_path: Path) -> None:
    created = _run(tmp_path, action="create", file="report.xlsx", sheet="Data")
    assert created.success is True
    assert created.metadata["output_file"] == "report.xlsx"

    written = _run(
        tmp_path,
        action="write_cells",
        file="report.xlsx",
        sheet="Data",
        cells={"A1": "标题", "B2": 123, "C2": 7},
    )
    assert written.success is True
    assert written.metadata["cells_written"] == 3

    grid = _run(
        tmp_path,
        action="write_cells",
        file="report.xlsx",
        sheet="Data",
        cells=[["Name", "Score"], ["Ada", 90]],
    )
    assert grid.success is True

    read = _run(tmp_path, action="read", file="report.xlsx", sheet="Data", read_range="A1:C2")
    assert read.success is True
    # 2D array write starts at A1, so A1/B1/A2/B2 come from the grid; C2 is kept.
    assert "Name" in read.content
    assert "Ada" in read.content
    assert "7" in read.content

    mapping_book = _run(tmp_path, action="create", file="mapped.xlsx", sheet="Data")
    assert mapping_book.success is True
    _run(
        tmp_path,
        action="write_cells",
        file="mapped.xlsx",
        sheet="Data",
        cells={"A1": "标题", "B2": 123},
    )
    mapped = _run(tmp_path, action="read", file="mapped.xlsx", sheet="Data", read_range="A1:B2")
    assert mapped.success is True
    assert "标题" in mapped.content
    assert "123" in mapped.content
    workbook = _xlsx(tmp_path / "mapped.xlsx")
    assert workbook["Data"]["A1"].value == "标题"
    assert workbook["Data"]["B2"].value == 123


def test_formula_round_trip_returns_formula_string(tmp_path: Path) -> None:
    _run(tmp_path, action="create", file="calc.xlsx")
    _run(tmp_path, action="write_cells", file="calc.xlsx", cells={"B2": 10, "C2": 5})
    result = _run(
        tmp_path,
        action="formula",
        file="calc.xlsx",
        formula_cells={"D2": "=SUM(B2:C2)"},
    )
    assert result.success is True
    assert result.metadata["formulas"] == 1

    read = _run(tmp_path, action="read", file="calc.xlsx", read_range="B2:D2")
    assert read.success is True
    assert "=SUM(B2:C2)" in read.content
    assert _xlsx(tmp_path / "calc.xlsx").active["D2"].value == "=SUM(B2:C2)"


def test_style_applies_without_error(tmp_path: Path) -> None:
    _run(tmp_path, action="create", file="styled.xlsx")
    _run(tmp_path, action="write_cells", file="styled.xlsx", cells={"A1": "H", "B1": "I"})
    result = _run(
        tmp_path,
        action="style",
        file="styled.xlsx",
        styles=[
            {
                "target": "A1:B1",
                "bold": True,
                "bg": "#B0501E",
                "color": "#FFFFFF",
                "font_size": 12,
            }
        ],
    )
    assert result.success is True
    cell = _xlsx(tmp_path / "styled.xlsx").active["A1"]
    assert cell.font.bold is True
    assert cell.font.size == 12
    rgb = str(getattr(cell.fill.fgColor, "rgb", cell.fill.fgColor) or "")
    assert rgb.upper().endswith("B0501E")


def test_style_rejects_empty_list(tmp_path: Path) -> None:
    _run(tmp_path, action="create", file="styled.xlsx")
    result = _run(tmp_path, action="style", file="styled.xlsx", styles=[])
    assert result.success is False
    assert "styles" in result.content.lower()


def test_chart_creates_bar_chart(tmp_path: Path) -> None:
    _run(tmp_path, action="create", file="chart.xlsx")
    _run(
        tmp_path,
        action="write_cells",
        file="chart.xlsx",
        cells=[["Month", "Sales"], ["Jan", 10], ["Feb", 20], ["Mar", 15]],
    )
    result = _run(
        tmp_path,
        action="chart",
        file="chart.xlsx",
        chart={"type": "bar", "data_range": "A1:B4", "title": "Sales"},
    )
    assert result.success is True
    charts = _xlsx(tmp_path / "chart.xlsx").active._charts
    assert len(charts) >= 1


def test_chart_rejects_unknown_type(tmp_path: Path) -> None:
    _run(tmp_path, action="create", file="chart.xlsx")
    _run(tmp_path, action="write_cells", file="chart.xlsx", cells=[["A", "B"], [1, 2]])
    result = _run(
        tmp_path,
        action="chart",
        file="chart.xlsx",
        chart={"type": "radar", "data_range": "A1:B2"},
    )
    assert result.success is False
    assert "bar" in result.content


def test_export_doc_renders_heading_table_list(tmp_path: Path) -> None:
    markdown = "\n".join(
        [
            "# Title",
            "Intro paragraph.",
            "## Section",
            "- alpha",
            "- beta",
            "| Col A | Col B |",
            "| --- | --- |",
            "| 1 | 2 |",
        ]
    )
    result = _run(tmp_path, action="export_doc", file="brief.docx", content=markdown)
    assert result.success is True
    path = tmp_path / "brief.docx"
    assert path.is_file()
    assert path.stat().st_size > 0

    from docx import Document

    document = Document(path)
    styles = [paragraph.style.name for paragraph in document.paragraphs if paragraph.style]
    assert any(name.startswith("Heading") for name in styles)
    assert any("List" in name for name in styles)
    assert len(document.tables) >= 1


def test_export_doc_requires_content(tmp_path: Path) -> None:
    result = _run(tmp_path, action="export_doc", file="brief.docx", content="   ")
    assert result.success is False
    assert "content" in result.content.lower()


def test_export_slide_builds_widescreen_deck(tmp_path: Path) -> None:
    outline = "\n".join(
        [
            "# Kickoff",
            "- goal one",
            "- goal two",
            "# Next steps",
            "- ship it",
        ]
    )
    result = _run(tmp_path, action="export_slide", file="deck.pptx", content=outline)
    assert result.success is True
    assert result.metadata["slides"] == 2
    path = tmp_path / "deck.pptx"
    assert path.is_file() and path.stat().st_size > 0

    from pptx import Presentation
    from pptx.util import Inches

    presentation = Presentation(path)
    assert len(presentation.slides) == 2
    assert abs(int(presentation.slide_width) - int(Inches(13.333))) < 2000
    assert presentation.slide_height == Inches(7.5)


def test_export_slide_requires_content(tmp_path: Path) -> None:
    result = _run(tmp_path, action="export_slide", file="deck.pptx")
    assert result.success is False
    assert "content" in result.content.lower()


def test_invalid_action_returns_failure(tmp_path: Path) -> None:
    result = _run(tmp_path, action="explode", file="x.xlsx")
    assert result.success is False
    assert "Invalid action" in result.content
    assert "create" in result.content


def test_write_cells_without_create_errors_clearly(tmp_path: Path) -> None:
    result = _run(tmp_path, action="write_cells", file="missing.xlsx", cells={"A1": 1})
    assert result.success is False
    assert "does not exist" in result.content
    assert "create" in result.content


def test_create_requires_file(tmp_path: Path) -> None:
    result = _run(tmp_path, action="create")
    assert result.success is False
    assert "file" in result.content.lower()


def test_path_escape_is_rejected(tmp_path: Path) -> None:
    result = _run(tmp_path, action="create", file="../escape.xlsx")
    assert result.success is False
    assert "outside" in result.content.lower()
    assert not (tmp_path.parent / "escape.xlsx").exists()


def test_add_sheet_and_screenshot_hint(tmp_path: Path) -> None:
    _run(tmp_path, action="create", file="book.xlsx")
    added = _run(tmp_path, action="add_sheet", file="book.xlsx", sheet="Extra")
    assert added.success is True
    _run(
        tmp_path,
        action="write_cells",
        file="book.xlsx",
        sheet="Extra",
        cells={"A1": "hello"},
    )
    hint = _run(tmp_path, action="screenshot_hint", file="book.xlsx", sheet="Extra")
    assert hint.success is True
    assert "hello" in hint.content
    preview = tmp_path / "book_preview.png"
    if preview.exists():
        assert preview.stat().st_size > 0


def test_execute_does_not_raise_on_failure() -> None:
    result = execute_office_document({"action": "not-a-real-action", "file": "x.xlsx"})
    assert isinstance(result, ToolResult)
    assert result.success is False


def test_office_document_is_registered_and_always_on() -> None:
    assert "office_document" in BUILTIN_TOOL_NAMES
    assert "office_document" in CONFIGURABLE_BUILTIN_TOOL_NAMES
    tools = compose_enabled_tools(
        registry=_EmptyRegistry(),
        requested_tools=[],
        optional_whitelist=[],
        mount_flags=ToolMountFlags(),
    )
    assert "office_document" in tools


def test_pipeline_injects_workspace_kwargs(monkeypatch, tmp_path: Path) -> None:
    from knorvia.agents.chat.agentic_pipeline import AgenticChatPipeline
    from knorvia.core.context import UnifiedContext
    from knorvia.services.path_service import PathService

    service = PathService(workspace_root=tmp_path)
    monkeypatch.setattr("knorvia.services.path_service.get_path_service", lambda: service)
    pipeline = AgenticChatPipeline.__new__(AgenticChatPipeline)
    monkeypatch.setattr(
        AgenticChatPipeline, "_current_user_id", lambda self: "u_ada", raising=False
    )
    context = UnifiedContext(
        session_id="s1", user_message="make a sheet", metadata={"turn_id": "t-1"}
    )
    kwargs = pipeline._augment_tool_kwargs("office_document", {"action": "create"}, context)
    assert kwargs["_sandbox_user_id"] == "u_ada"
    workdir = Path(kwargs["_sandbox_workdir"])
    assert workdir.name == "exec"
    assert workdir.is_dir()
    assert kwargs["_workspace_dir"] == kwargs["_sandbox_workdir"]
    assert service.is_public_output_path(workdir / "report.xlsx") is False
    (workdir / "report.xlsx").write_bytes(b"PK\x03\x04placeholder")
    assert service.is_public_output_path(workdir / "report.xlsx")

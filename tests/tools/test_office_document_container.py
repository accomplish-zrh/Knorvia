"""Coverage for ``office_document`` .univer container actions."""

from __future__ import annotations

from pathlib import Path
from typing import Any

from knorvia.core.tool_protocol import ToolResult
from knorvia.services.univer_container import list_units, open_univer, read_unit_bytes
from knorvia.tools.office_document import ACTIONS, execute_office_document


def _run(workspace: Path, **kwargs: Any) -> ToolResult:
    return execute_office_document(
        {
            "_workspace_dir": str(workspace),
            "_sandbox_workdir": str(workspace),
            "as_draft": False,
            **kwargs,
        }
    )


def test_container_actions_are_registered() -> None:
    assert "container_new" in ACTIONS
    assert "container_add" in ACTIONS
    assert "container_export_unit" in ACTIONS


def test_container_new_add_export_round_trip(tmp_path: Path) -> None:
    created = _run(tmp_path, action="create", file="sales.xlsx", sheet="Data")
    assert created.success is True
    _run(
        tmp_path,
        action="write_cells",
        file="sales.xlsx",
        sheet="Data",
        cells={"A1": "Item", "B1": "Qty", "A2": "Widgets", "B2": 3},
    )

    bundle = _run(
        tmp_path,
        action="container_new",
        file="pack.univer",
        source_file="sales.xlsx",
        unit_id="sheet",
        unit_name="Sales",
    )
    assert bundle.success is True
    assert (tmp_path / "pack.univer").is_file()
    assert len(bundle.metadata["units"]) == 1

    added = _run(
        tmp_path,
        action="container_add",
        file="pack.univer",
        unit_type="slide",
        unit_id="slide",
        unit_name="Pitch",
        refs=[{"to": "sheet", "range": "A1:B2"}],
    )
    assert added.success is True
    manifest = open_univer(tmp_path / "pack.univer")
    assert [unit["id"] for unit in manifest["units"]] == ["sheet", "slide"]
    assert manifest["refs"][0]["to"] == "sheet"
    payload, _ = read_unit_bytes(tmp_path / "pack.univer", "slide")
    from io import BytesIO

    from pptx import Presentation

    slide = Presentation(BytesIO(payload)).slides[0]
    texts = [shape.text_frame.text for shape in slide.shapes if shape.has_text_frame]
    assert any("数据来源: sheet!A1:B2" in text for text in texts)

    doc = _run(
        tmp_path,
        action="container_add",
        file="pack.univer",
        unit_type="doc",
        unit_id="doc",
        unit_name="Notes",
    )
    assert doc.success is True
    assert len(list_units(tmp_path / "pack.univer")) == 3

    exported = _run(
        tmp_path,
        action="container_export_unit",
        file="pack.univer",
        unit_id="sheet",
        export_file="exported_sales",
    )
    assert exported.success is True
    assert (tmp_path / "exported_sales.xlsx").is_file()


def test_container_add_without_new_fails(tmp_path: Path) -> None:
    result = _run(
        tmp_path,
        action="container_add",
        file="missing.univer",
        unit_type="sheet",
    )
    assert result.success is False
    assert "container_new" in result.content


def test_container_export_requires_unit_id(tmp_path: Path) -> None:
    _run(tmp_path, action="container_new", file="empty.univer")
    result = _run(
        tmp_path,
        action="container_export_unit",
        file="empty.univer",
    )
    assert result.success is False
    assert "unit_id" in result.content


def test_invalid_unit_type_fails(tmp_path: Path) -> None:
    _run(tmp_path, action="container_new", file="box.univer")
    result = _run(
        tmp_path,
        action="container_add",
        file="box.univer",
        unit_type="pdf",
    )
    assert result.success is False
    assert "unit_type" in result.content

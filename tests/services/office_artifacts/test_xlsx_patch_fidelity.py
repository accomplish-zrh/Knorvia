"""Format-fidelity of the narrow OOXML patcher for *imported* workbooks.

These files came from elsewhere, so a patch may only change what the operation
asked for: clearing a value must not erase its formatting, resizing one column
inside a grouped ``<col>`` range must not strip the range's other properties,
and copying a style onto blank cells must actually create those cells.
"""

from __future__ import annotations

import hashlib
import io
from pathlib import Path
import zipfile

from lxml import etree
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
import pytest

from knorvia.services.office_artifacts.service import OfficeArtifactService

SHEET_PART = "xl/worksheets/sheet1.xml"
MAIN_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
NS = {"m": MAIN_NS}


def _rebuild_zip(data: bytes, replacements: dict[str, bytes]) -> bytes:
    with zipfile.ZipFile(io.BytesIO(data)) as bundle:
        names = bundle.namelist()
        entries = {name: bundle.read(name) for name in names}
    entries.update(replacements)
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as bundle:
        for name in names:
            bundle.writestr(name, entries[name])
    return out.getvalue()


def _styled_workbook() -> tuple[bytes, str]:
    """An Excel-shaped fixture: a styled cell and one grouped A:E column range.

    openpyxl cannot emit a grouped ``<col>`` carrying width, style and hidden
    together, so the range is injected into the package the way a real Excel
    file would have it.
    """
    book = Workbook()
    ws = book.active
    assert ws is not None
    ws.title = "Data"
    ws["A1"] = "keep my format"
    ws["A1"].font = Font(bold=True, color="FF0000")
    ws["A1"].fill = PatternFill("solid", fgColor="FFFF00")
    ws["A3"] = "row three"
    buffer = io.BytesIO()
    book.save(buffer)
    data = buffer.getvalue()

    root = etree.fromstring(_part(data, SHEET_PART))
    # Reference a cellXfs index that really exists in this package.
    col_style = (_cell_xml(data, "A1") or etree.Element("c")).get("s") or "0"
    cols = etree.Element(f"{{{MAIN_NS}}}cols")
    etree.SubElement(
        cols,
        f"{{{MAIN_NS}}}col",
        {
            "min": "1",
            "max": "5",
            "width": "22.5",
            "style": col_style,
            "hidden": "1",
            "customWidth": "1",
        },
    )
    root.find(f"{{{MAIN_NS}}}sheetData").addprevious(cols)
    return _rebuild_zip(
        data,
        {SHEET_PART: etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)},
    ), col_style


def _imported_service(
    tmp_path: Path, data: bytes
) -> tuple[OfficeArtifactService, str, str, str]:
    service = OfficeArtifactService(task_dir=tmp_path, workspace_dir=tmp_path / "exec")
    draft_id = service.store.create_draft()
    artifact_id = service.store.add_artifact(
        draft_id,
        filename="book.xlsx",
        mime="",
        kind="xlsx",
        origin_kind="attachment",
        origin_ref="attachment:book.xlsx",
        data=data,
        origin_base_hash=hashlib.sha256(data).hexdigest(),
    )
    sheet = load_workbook(io.BytesIO(data)).sheetnames[0]
    return service, draft_id, artifact_id, sheet


def _apply(
    service: OfficeArtifactService,
    draft_id: str,
    artifact_id: str,
    operations: list[dict[str, object]],
    base: int = 0,
) -> dict:
    return service.apply_operations(draft_id, artifact_id, base, operations, actor="tester")


def _part(data: bytes, name: str) -> bytes:
    with zipfile.ZipFile(io.BytesIO(data)) as bundle:
        return bundle.read(name)


def _cell_xml(data: bytes, ref: str) -> etree._Element | None:
    root = etree.fromstring(_part(data, SHEET_PART))
    for cell in root.iter(f"{{{MAIN_NS}}}c"):
        if (cell.get("r") or "").upper() == ref:
            return cell
    return None


def _row_xml(data: bytes, row: int) -> etree._Element | None:
    root = etree.fromstring(_part(data, SHEET_PART))
    for element in root.iter(f"{{{MAIN_NS}}}row"):
        if element.get("r") == str(row):
            return element
    return None


def _cols(data: bytes) -> list[dict[str, str]]:
    root = etree.fromstring(_part(data, SHEET_PART))
    container = root.find(f"{{{MAIN_NS}}}cols")
    if container is None:
        return []
    return [dict(element.attrib) for element in container.findall(f"{{{MAIN_NS}}}col")]


def test_clear_cells_keeps_the_cells_formatting(tmp_path: Path) -> None:
    source, col_style = _styled_workbook()
    before = _cell_xml(source, "A1")
    assert before is not None
    style_id = before.get("s")
    assert style_id

    service, draft_id, artifact_id, sheet = _imported_service(tmp_path, source)
    _apply(
        service,
        draft_id,
        artifact_id,
        [{"op": "clear_cells", "sheet": sheet, "range": "A1"}],
    )

    after = _cell_xml(service.store.current_bytes(draft_id, artifact_id), "A1")
    assert after is not None, "clear_cells deleted the <c> element and its style with it"
    assert after.get("s") == style_id
    assert after.find(f"{{{MAIN_NS}}}v") is None
    workbook = load_workbook(io.BytesIO(service.store.current_bytes(draft_id, artifact_id)))
    cell = workbook[sheet]["A1"]
    assert cell.value is None
    assert cell.fill.start_color.rgb == "00FFFF00"
    assert cell.font.bold is True


def test_clear_cells_does_not_drop_row_attributes(tmp_path: Path) -> None:
    source, col_style = _styled_workbook()
    service, draft_id, artifact_id, sheet = _imported_service(tmp_path, source)
    _apply(
        service,
        draft_id,
        artifact_id,
        [{"op": "clear_cells", "sheet": sheet, "range": "A3"}],
    )
    row = _row_xml(service.store.current_bytes(draft_id, artifact_id), 3)
    assert row is not None, "clearing the row's only cell removed the <row> element"


def test_resizing_one_column_in_a_grouped_range_keeps_every_other_property(
    tmp_path: Path,
) -> None:
    source, col_style = _styled_workbook()
    grouped = _cols(source)
    assert len(grouped) == 1
    assert grouped[0]["min"] == "1" and grouped[0]["max"] == "5"
    assert grouped[0]["width"] == "22.5" and grouped[0]["hidden"] == "1"
    assert grouped[0]["style"] == col_style

    service, draft_id, artifact_id, sheet = _imported_service(tmp_path, source)
    _apply(
        service,
        draft_id,
        artifact_id,
        [{"op": "set_column_width", "sheet": sheet, "column": "C", "width": 8.0}],
    )

    cols = _cols(service.store.current_bytes(draft_id, artifact_id))
    by_span = {(c["min"], c["max"]): c for c in cols}
    assert [(c["min"], c["max"]) for c in cols] == [("1", "2"), ("3", "3"), ("4", "5")]
    for span in (("1", "2"), ("4", "5")):
        kept = by_span[span]
        assert kept["width"] == "22.5"
        assert kept.get("hidden") == "1"
        assert kept.get("style") == col_style
    assert by_span[("3", "3")]["width"] == "8.0"


def test_copy_style_onto_blank_cells_actually_writes_the_style(tmp_path: Path) -> None:
    source, col_style = _styled_workbook()
    service, draft_id, artifact_id, sheet = _imported_service(tmp_path, source)
    source_cell = _cell_xml(source, "A1")
    assert source_cell is not None
    style_id = source_cell.get("s")

    result = _apply(
        service,
        draft_id,
        artifact_id,
        [{"op": "copy_style", "sheet": sheet, "source_cell": "A1", "target_range": "C5:D6"}],
    )
    current = service.store.current_bytes(draft_id, artifact_id)

    assert result["touched_cells"] == 4
    for ref in ("C5", "D5", "C6", "D6"):
        cell = _cell_xml(current, ref)
        assert cell is not None, f"{ref} was skipped, so the reported change never happened"
        assert cell.get("s") == style_id


def test_patched_workbook_still_reopens_in_openpyxl(tmp_path: Path) -> None:
    source, col_style = _styled_workbook()
    service, draft_id, artifact_id, sheet = _imported_service(tmp_path, source)
    _apply(
        service,
        draft_id,
        artifact_id,
        [
            {"op": "set_cell", "sheet": sheet, "cell": "B2", "text": "hello"},
            {"op": "clear_cells", "sheet": sheet, "range": "A1"},
            {"op": "set_column_width", "sheet": sheet, "column": "C", "width": 12.0},
        ],
    )
    current = service.store.current_bytes(draft_id, artifact_id)
    workbook = load_workbook(io.BytesIO(current))
    assert workbook[sheet]["B2"].value == "hello"
    assert workbook[sheet].column_dimensions["C"].width == 12.0

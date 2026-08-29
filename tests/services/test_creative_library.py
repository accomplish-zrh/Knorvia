from __future__ import annotations

from io import BytesIO
import json

from PIL import Image
import pytest

from knorvia.services.creative_library.canvas import add_canvas_node, normalize_library_canvas
from knorvia.services.creative_library.office import encode_library_office, extract_library_office
from knorvia.services.creative_library.store import CreativeLibraryStore


def _png() -> bytes:
    buffer = BytesIO()
    Image.new("RGB", (8, 8), (12, 80, 160)).save(buffer, format="PNG")
    return buffer.getvalue()


def test_text_asset_round_trip(tmp_path) -> None:
    store = CreativeLibraryStore(tmp_path)
    asset = store.create_text_asset(title="Wardrobe", content="navy coat, white shirt")
    listed = store.list_assets(keyword="navy")
    assert listed["total"] == 1
    assert listed["items"][0]["id"] == asset["id"]
    updated = store.update_asset(asset["id"], {"note": "keep this look"})
    assert updated["note"] == "keep this look"
    assert store.delete_asset(asset["id"]) is True
    assert store.get_asset(asset["id"]) is None


def test_image_asset_bytes(tmp_path) -> None:
    store = CreativeLibraryStore(tmp_path)
    asset = store.create_media_asset(_png(), "image/png", title="Hero")
    payload = store.asset_bytes(asset["id"])
    assert payload is not None
    data, mime = payload
    assert mime == "image/png"
    assert data[:8] == b"\x89PNG\r\n\x1a\n"


def test_builtin_prompts_seeded_and_user_prompts_editable(tmp_path) -> None:
    store = CreativeLibraryStore(tmp_path)
    built = store.list_prompts(origin="builtin")
    assert built["total"] >= 6
    prompt = store.create_prompt(title="Mine", body="soft window light", language="en")
    assert prompt["origin"] == "user"
    changed = store.update_prompt(prompt["id"], {"body": "softer window light"})
    assert changed["body"] == "softer window light"
    with pytest.raises(PermissionError):
        store.update_prompt(built["items"][0]["id"], {"body": "nope"})


def test_library_tree_personal_only_and_formats(tmp_path) -> None:
    store = CreativeLibraryStore(tmp_path)
    folder = store.create_entry(kind="folder", title="Notes")
    md = store.create_entry(
        kind="markdown",
        title="Brief",
        parent_id=folder["id"],
        content="# Hello",
    )
    store.create_entry(kind="csv", title="Rows", parent_id=folder["id"], content="a,b\n1,2")
    canvas = store.create_entry(kind="canvas", title="Board")
    tree = store.list_tree()
    titles = [item["title"] for item in tree["items"]]
    assert "Notes" in titles
    assert "Board" in titles
    assert all("team" not in str(item.get("title") or "").lower() for item in tree["items"])
    notes = next(item for item in tree["items"] if item["id"] == folder["id"])
    assert [child["kind"] for child in notes["children"]] == ["csv", "markdown"] or {
        child["kind"] for child in notes["children"]
    } == {"csv", "markdown"}
    loaded = store.get_entry(md["id"])
    assert loaded is not None
    assert loaded["content"] == "# Hello"
    changed = store.update_entry(md["id"], {"content": "# Edited"})
    assert changed["content"] == "# Edited"
    canvas_doc = store.get_entry(canvas["id"])
    assert canvas_doc is not None
    assert '"nodes"' in canvas_doc["content"]
    assert store.delete_entry(folder["id"]) is True
    assert store.get_entry(md["id"]) is None


def test_library_imports_legacy_assets_once(tmp_path) -> None:
    store = CreativeLibraryStore(tmp_path)
    store.create_text_asset(title="Old note", content="keep me")
    first = store.list_tree()
    second = store.list_tree()
    imported = [item for item in first["items"] if item["title"] == "Imported assets"]
    assert len(imported) == 1
    assert any(child["title"] == "Old note" for child in imported[0]["children"])
    assert second["total"] == first["total"]


def test_library_word_and_excel_create_extract_update(tmp_path) -> None:
    store = CreativeLibraryStore(tmp_path)
    word = store.create_entry(kind="word", title="Brief", content="word-body-v1")
    excel = store.create_entry(kind="excel", title="Rows", content="excel-body-v1")
    word_bytes = store.entry_bytes(word["id"])
    excel_bytes = store.entry_bytes(excel["id"])
    assert word_bytes is not None
    assert excel_bytes is not None
    assert word_bytes[0].startswith(b"PK")
    assert excel_bytes[0].startswith(b"PK")
    assert word_bytes[1].endswith("wordprocessingml.document")
    assert excel_bytes[1].endswith("spreadsheetml.sheet")
    assert extract_library_office("word", word_bytes[0]) == "word-body-v1"
    assert extract_library_office("excel", excel_bytes[0]) == "excel-body-v1"
    assert word["content"] == "word-body-v1"
    assert excel["content"] == "excel-body-v1"
    word_path = tmp_path / "files" / f"{word['id']}.docx"
    excel_path = tmp_path / "files" / f"{excel['id']}.xlsx"
    assert word_path.is_file()
    assert excel_path.is_file()
    updated_word = store.update_entry(word["id"], {"content": "word-body-v2"})
    updated_excel = store.update_entry(excel["id"], {"content": "excel-body-v2"})
    word_again = store.entry_bytes(word["id"])
    excel_again = store.entry_bytes(excel["id"])
    assert word_again is not None and excel_again is not None
    assert extract_library_office("word", word_again[0]) == "word-body-v2"
    assert extract_library_office("excel", excel_again[0]) == "excel-body-v2"
    assert updated_word["content"] == "word-body-v2"
    assert updated_excel["content"] == "excel-body-v2"
    uploaded = store.upload_entry(
        encode_library_office("word", "uploaded-word"),
        "note.docx",
        title="Upload",
    )
    assert uploaded["kind"] == "word"
    assert uploaded["content"] == "uploaded-word"


def test_library_excel_replace_bytes_keeps_other_sheets(tmp_path) -> None:
    from openpyxl import Workbook, load_workbook

    store = CreativeLibraryStore(tmp_path)
    workbook = Workbook()
    sales = workbook.active
    sales.title = "Sales"
    sales["A1"] = "header"
    sales["B1"] = 10
    sales["B2"] = "=B1*2"
    notes = workbook.create_sheet("Notes")
    notes["A1"] = "second-sheet"
    buffer = BytesIO()
    workbook.save(buffer)
    payload = buffer.getvalue()

    uploaded = store.upload_entry(payload, "book.xlsx", title="Book")
    assert uploaded["kind"] == "excel"
    replaced = store.replace_entry_bytes(uploaded["id"], payload)
    data = store.entry_bytes(replaced["id"])
    assert data is not None
    loaded = load_workbook(BytesIO(data[0]))
    assert loaded.sheetnames == ["Sales", "Notes"]
    assert loaded["Notes"]["A1"].value == "second-sheet"
    assert loaded["Sales"]["A1"].value == "header"
    assert loaded["Sales"]["B2"].value == "=B1*2"

    word = store.create_entry(kind="word", title="Memo", content="plain")
    with pytest.raises(ValueError, match="Only Excel"):
        store.replace_entry_bytes(word["id"], payload)
    with pytest.raises(ValueError, match="Not a valid Excel"):
        store.replace_entry_bytes(uploaded["id"], b"not-xlsx")

    flattened = store.update_entry(uploaded["id"], {"content": "flat-text"})
    destroyed = load_workbook(BytesIO(store.entry_bytes(flattened["id"])[0]))
    assert destroyed.sheetnames == ["Sheet"]
    assert extract_library_office("excel", store.entry_bytes(flattened["id"])[0]) == "flat-text"


def test_library_canvas_add_card_persists_on_entry(tmp_path) -> None:
    store = CreativeLibraryStore(tmp_path)
    created = store.create_entry(kind="canvas", title="Board")
    document = add_canvas_node(
        normalize_library_canvas(json.loads(created["content"] or "{}")),
        title="Card",
        text="pinned-card",
    )
    store.update_entry(created["id"], {"content": json.dumps(document)})
    loaded = store.get_entry(created["id"])
    assert loaded is not None
    parsed = json.loads(loaded["content"])
    assert parsed["nodes"][0]["text"] == "pinned-card"
    assert parsed["nodes"][0]["title"] == "Card"


def test_conversation_messages(tmp_path) -> None:
    store = CreativeLibraryStore(tmp_path)
    conversation = store.create_conversation("Lookbook")
    store.add_message(conversation["id"], role="user", content="make a poster")
    store.add_message(
        conversation["id"],
        role="assistant",
        content="queued",
        brief={"user_prompt": "make a poster"},
        job={"studio": "image", "job_id": "job_1"},
    )
    messages = store.list_messages(conversation["id"])
    assert len(messages) == 2
    assert messages[1]["job"]["job_id"] == "job_1"
    assert "user_prompt" in messages[1]["brief"]

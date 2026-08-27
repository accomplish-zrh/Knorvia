"""Unit tests for the ZIP-backed ``.univer`` multi-unit container."""

from __future__ import annotations

from pathlib import Path

import pytest

from knorvia.services.univer_container import (
    ContainerError,
    add_unit,
    export_unit,
    list_units,
    open_univer,
    pack_univer,
    read_unit_bytes,
    remove_unit,
)


def test_pack_open_round_trip(tmp_path: Path) -> None:
    out = tmp_path / "pack.univer"
    manifest = pack_univer(
        [
            {"id": "sheet", "type": "sheet", "name": "Data"},
            {"id": "doc", "type": "doc", "name": "Notes"},
        ],
        out,
    )
    assert out.is_file()
    assert out.stat().st_size > 0
    assert [unit["id"] for unit in manifest["units"]] == ["sheet", "doc"]

    loaded = open_univer(out)
    assert loaded["version"] == 1
    assert len(loaded["units"]) == 2
    assert list_units(out)[0]["file"] == "units/sheet.xlsx"


def test_add_unit_records_refs_and_slide_placeholder(tmp_path: Path) -> None:
    out = tmp_path / "deck.univer"
    pack_univer([{"id": "sheet", "type": "sheet", "name": "Sales"}], out)
    unit = add_unit(
        out,
        "slide",
        "Pitch",
        unit_id="slide",
        refs=[{"to": "sheet", "range": "A1:B5"}],
    )
    assert unit["id"] == "slide"
    manifest = open_univer(out)
    assert manifest["refs"] == [
        {
            "from": "slide",
            "to": "sheet",
            "range": "A1:B5",
            "kind": "data_source",
        }
    ]
    payload, _entry = read_unit_bytes(out, "slide")
    from io import BytesIO

    from pptx import Presentation

    slide = Presentation(BytesIO(payload)).slides[0]
    texts = [shape.text_frame.text for shape in slide.shapes if shape.has_text_frame]
    assert any("数据来源: sheet!A1:B5" in text for text in texts)


def test_export_and_remove_unit(tmp_path: Path) -> None:
    out = tmp_path / "bundle.univer"
    pack_univer(
        [
            {"id": "sheet", "type": "sheet"},
            {"id": "doc", "type": "doc"},
        ],
        out,
    )
    exported = export_unit(out, "sheet", tmp_path / "out" / "data")
    assert exported.suffix == ".xlsx"
    assert exported.is_file() and exported.stat().st_size > 0

    removed = remove_unit(out, "doc")
    assert removed["id"] == "doc"
    assert [unit["id"] for unit in list_units(out)] == ["sheet"]


def test_corrupt_zip_raises(tmp_path: Path) -> None:
    bad = tmp_path / "broken.univer"
    bad.write_bytes(b"not-a-zip")
    with pytest.raises(ContainerError, match="corrupt"):
        open_univer(bad)


def test_missing_unit_and_duplicate_id(tmp_path: Path) -> None:
    out = tmp_path / "one.univer"
    pack_univer([{"id": "sheet", "type": "sheet"}], out)
    with pytest.raises(ContainerError, match="not found"):
        read_unit_bytes(out, "missing")
    with pytest.raises(ContainerError, match="already exists"):
        add_unit(out, "sheet", unit_id="sheet")


def test_pack_from_source_bytes(tmp_path: Path) -> None:
    source = tmp_path / "seed.xlsx"
    pack_univer([{"id": "tmp", "type": "sheet"}], tmp_path / "seed.univer")
    # Reuse an exported native file as the source for a fresh container.
    export_unit(tmp_path / "seed.univer", "tmp", source)
    out = tmp_path / "from-source.univer"
    manifest = pack_univer(
        [{"id": "sheet", "type": "sheet", "source": source}],
        out,
    )
    assert manifest["units"][0]["file"] == "units/sheet.xlsx"
    payload, _ = read_unit_bytes(out, "sheet")
    assert payload == source.read_bytes()
